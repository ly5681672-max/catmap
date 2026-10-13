import importlib.util
import csv
import tempfile
import unittest
from pathlib import Path


SCRIPT = Path(__file__).parents[1] / "scripts" / "extract_neb_gibbs.py"
spec = importlib.util.spec_from_file_location("extract_neb_gibbs", SCRIPT)
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)


class WorkbookExtractionTests(unittest.TestCase):
    def test_real_workbook_contains_all_fourteen_requested_systems(self):
        import openpyxl

        path = Path(r"H:\catmap\catmap\lyq_alkene_epoxidation\己烯环氧化.xlsx")
        self.assertTrue(path.exists())
        wb = openpyxl.load_workbook(path, data_only=False, read_only=True)
        groups = module.catalyst_groups(wb["吉布斯自由能汇总"])
        self.assertEqual(set(module.CATALYST_ORDER), set(module.CATALYST_ORDER) & set(groups))
        self.assertIn("tinb", groups)
        self.assertNotIn("tinb", module.CATALYST_ORDER)

    def test_tire_neb_state_rows_and_formula_cache_match_reference_targets(self):
        import openpyxl

        path = Path(r"H:\catmap\catmap\lyq_alkene_epoxidation\己烯环氧化.xlsx")
        formulas = openpyxl.load_workbook(path, data_only=False, read_only=True)["吉布斯自由能汇总"]
        values = openpyxl.load_workbook(path, data_only=True, read_only=True)["吉布斯自由能汇总"]
        start = module.catalyst_groups(formulas)["tire"]
        expected_rows = {3: 12, 4: 16, 5: 20, 6: 24, 7: 28, 8: 32}
        for state_index, row in expected_rows.items():
            self.assertEqual(formulas.cell(row, start).value, state_index)
        neb1_is = values.cell(12, start + 5).value
        neb1_ts = values.cell(16, start + 5).value
        neb1_fs = values.cell(20, start + 5).value
        neb2_is = values.cell(24, start + 5).value
        neb2_ts = values.cell(28, start + 5).value
        neb2_fs = values.cell(32, start + 5).value
        self.assertAlmostEqual(neb1_ts - neb1_is, 0.444846, places=6)
        self.assertAlmostEqual(neb1_ts - neb1_fs, 2.372278, places=6)
        self.assertAlmostEqual(neb2_ts - neb2_is, 0.217873, places=6)
        self.assertAlmostEqual(neb2_ts - neb2_fs, 1.030708, places=6)

    def test_formula_audit_preserves_formula_and_cached_value_separately(self):
        import openpyxl

        path = Path(r"H:\catmap\catmap\lyq_alkene_epoxidation\己烯环氧化.xlsx")
        formulas = openpyxl.load_workbook(path, data_only=False, read_only=True)["吉布斯自由能汇总"]
        values = openpyxl.load_workbook(path, data_only=True, read_only=True)["吉布斯自由能汇总"]
        self.assertTrue(str(formulas["AC16"].value).startswith("="))
        self.assertAlmostEqual(values["AC16"].value - values["AC12"].value, 0.444846, places=6)
        self.assertAlmostEqual(values["AD16"].value, 0.444846, places=6)

    def test_extraction_outputs_six_neb_anchor_states_per_requested_catalyst(self):
        workbook = Path(r"H:\catmap\catmap\lyq_alkene_epoxidation\己烯环氧化.xlsx")
        with tempfile.TemporaryDirectory() as tmp:
            module.extract_workbook(workbook, Path(tmp))
            with open(Path(tmp) / "original_neb_state_inventory.csv", encoding="utf-8-sig") as stream:
                rows = list(csv.DictReader(stream))
            self.assertEqual(len(rows), 14 * 6)
            tire_neb1_is = next(r for r in rows if r["catalyst"] == "tire" and r["NEB_role"] == "NEB1_IS")
            self.assertEqual(tire_neb1_is["raw_G_cell"], "AC12")
            timn_neb2_fs = next(r for r in rows if r["catalyst"] == "timn" and r["NEB_role"] == "NEB2_FS")
            self.assertEqual(timn_neb2_fs["raw_G_cell"], "N32")


if __name__ == "__main__":
    unittest.main()
