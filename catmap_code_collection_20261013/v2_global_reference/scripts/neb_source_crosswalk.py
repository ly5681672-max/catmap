"""Build a strict, auditable NEB-source crosswalk and Gibbs state table."""

from __future__ import annotations

import csv
import hashlib
import json
import math
import re
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook


V2_ROOT = Path(__file__).resolve().parents[1]
ROOT = V2_ROOT.parent
V1_DATA = ROOT / "data"
DATA = V2_ROOT / "data"
SOURCE_WORKBOOK = Path(r"H:\catmap\catmap\lyq_alkene_epoxidation\己烯环氧化.xlsx")
DFT_SOURCE = Path(r"H:\catmap\catmap\lyq_alkene_epoxidation\2-creating_microkinetic_model\DFT_TS_relative.csv")
REFERENCE_FORMULA = "G_full(X*) - G_clean_slab - sum_i(n_i * mu_i_ref)"


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1 << 20), b""):
            digest.update(block)
    return digest.hexdigest().upper()


def _require_exactly_one(matches: list, description: str):
    if len(matches) != 1:
        raise ValueError(f"{description}: expected exactly one source match, found {len(matches)}")
    return matches[0]


def _unique_index(rows: list[dict[str, str]], key_fn, description: str) -> dict:
    index = {}
    for row in rows:
        key = key_fn(row)
        if key in index:
            raise ValueError(f"duplicate source key in {description}: {key}")
        index[key] = row
    return index


def _source_key(catalyst: str, workbook: str, worksheet: str, source_row: str) -> tuple:
    try:
        row_number = int(source_row)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"invalid workbook source row {source_row!r} for {catalyst}/{worksheet}") from exc
    return catalyst.strip(), str(Path(workbook)), worksheet.strip(), row_number


def _same_number(actual, expected, tolerance: float = 1e-9) -> bool:
    try:
        return math.isclose(float(actual), float(expected), rel_tol=0.0, abs_tol=tolerance)
    except (TypeError, ValueError):
        return str(actual or "") == str(expected or "")


def _check_value(actual, expected, label: str, source: str) -> None:
    if not _same_number(actual, expected):
        raise ValueError(f"{label} mismatch at {source}: workbook={actual!r}; inventory={expected!r}")


def _gas_terms(*formulas: str) -> list[str]:
    combined = " ".join(formulas)
    terms = []
    if "F59" in combined:
        terms.append("H2O_g:F59")
    if "F60" in combined:
        terms.append("C6H12O_g:F60")
    return terms


def _mapping_fields(path: str, state: str) -> tuple[str, str, str, str]:
    if state == "IS":
        return "IS_source_cell", "IS_source_expression", "G_IS_workbook_or_normalized", "G_IS1" if path == "NEB1" else "G_IS2"
    if state == "TS":
        return "TS_source_cell", "TS_source_expression", "G_TS_workbook_or_normalized", "G_TS1" if path == "NEB1" else "G_TS2"
    return ("FS_source_cell", "FS_source_expression_or_adjustment", "G_FS_workbook_or_normalized",
            "G_FS1_total" if path == "NEB1" else "G_FS2_adjusted")


def build_rows() -> tuple[list[dict[str, str]], list[dict[str, str]], dict]:
    generic = _read_csv(V1_DATA / "original_gibbs_state_inventory.csv")
    neb = _read_csv(V1_DATA / "original_neb_state_inventory.csv")
    formula_audit = _read_csv(V1_DATA / "excel_gibbs_formula_audit.csv")
    path_mappings = _read_csv(V1_DATA / "neb_to_catmap_energy_mapping.csv")
    surfaces = _read_csv(V1_DATA / "SERIAL_14_SURFACE_MAPPING.csv")
    source_manifest = json.loads((V1_DATA / "source_hashes_before.json").read_text(encoding="utf-8"))
    hashes = source_manifest["sha256"]
    workbook_key, dft_key = str(SOURCE_WORKBOOK), str(DFT_SOURCE)
    if workbook_key not in hashes or dft_key not in hashes:
        raise ValueError("frozen SHA256 manifest must include the workbook and DFT_TS_relative.csv")
    workbook_hash, dft_hash = _sha256(SOURCE_WORKBOOK), _sha256(DFT_SOURCE)
    if workbook_hash != hashes[workbook_key]:
        raise ValueError(f"workbook SHA256 mismatch: {workbook_hash} != {hashes[workbook_key]}")
    if dft_hash != hashes[dft_key]:
        raise ValueError(f"DFT source SHA256 mismatch: {dft_hash} != {hashes[dft_key]}")

    target_catalysts = {row["catalyst"] for row in surfaces}
    if len(target_catalysts) != 14:
        raise ValueError(f"expected 14 V1 catalysts, found {len(target_catalysts)}")
    dft_rows = [row for row in _read_csv(DFT_SOURCE) if row["surface"] in target_catalysts]
    dft_by_catalyst = _unique_index(dft_rows, lambda row: row["surface"], "DFT catalyst rows")
    if set(dft_by_catalyst) != target_catalysts:
        raise ValueError(f"DFT source does not cover the exact V1 catalyst set: {sorted(target_catalysts ^ set(dft_by_catalyst))}")

    generic_by_source = _unique_index(
        generic,
        lambda row: _source_key(row["catalyst"], workbook_key, row["worksheet"], row["source_row"]),
        "generic Gibbs inventory",
    )
    neb_by_source = _unique_index(
        neb,
        lambda row: _source_key(row["catalyst"], workbook_key, row["worksheet"],
                                 row["source_row_from_DFT_TS_relative_reference_note"]),
        "NEB anchor inventory",
    )
    if len(neb_by_source) != 84:
        raise ValueError(f"expected 84 unique NEB source rows, found {len(neb_by_source)}")
    audit_by_cell: dict[tuple[str, str], list[dict[str, str]]] = {}
    for row in formula_audit:
        audit_by_cell.setdefault((row["worksheet"], row["cell"]), []).append(row)
    path_by_key = _unique_index(path_mappings, lambda row: (row["catalyst"], row["path"]), "NEB path mapping")
    expected_paths = {(catalyst, path) for catalyst in target_catalysts for path in ("NEB1", "NEB2")}
    if set(path_by_key) != expected_paths:
        raise ValueError(f"NEB path mapping set mismatch: missing={sorted(expected_paths-set(path_by_key))}, extra={sorted(set(path_by_key)-expected_paths)}")

    formula_wb = load_workbook(SOURCE_WORKBOOK, read_only=True, data_only=False)
    cached_wb = load_workbook(SOURCE_WORKBOOK, read_only=True, data_only=True)
    global_rows: list[dict[str, str]] = []
    crosswalk: list[dict[str, str]] = []
    mapped_paths: dict[tuple[str, str], dict[str, Decimal]] = {}
    anchor_crosswalk: dict[tuple[str, str], dict[str, str]] = {}

    def cell_values(sheet: str, cell: str) -> tuple[str, str]:
        if not cell:
            return "", ""
        return str(formula_wb[sheet][cell].value or ""), str(cached_wb[sheet][cell].value or "")

    try:
        for key in sorted(set(generic_by_source) | set(neb_by_source)):
            catalyst, workbook, sheet, source_row = key
            anchor = neb_by_source.get(key)
            state = generic_by_source.get(key)
            if anchor:
                role = anchor["NEB_role"]
                role_match = re.fullmatch(r"(NEB1|NEB2)_(IS|TS|FS)", role)
                if not role_match:
                    raise ValueError(f"invalid NEB role {role!r} at {key}")
                path, step_state = role_match.groups()
                path_record = path_by_key[(catalyst, path)]
                map_cell_field, map_formula_field, map_energy_field, dft_field = _mapping_fields(path, step_state)
                mapped_cell = path_record[map_cell_field]
                mapped_expression = path_record[map_formula_field]
                mapped_energy = path_record[map_energy_field]
                raw_cell = anchor["raw_G_cell"]
                source_row_match = re.search(r"(\d+)$", raw_cell)
                if not source_row_match or int(source_row_match.group(1)) != source_row:
                    raise ValueError(f"source row/cell conflict for {key}: {raw_cell}")
                raw_formula, raw_cached = cell_values(sheet, raw_cell)
                raw_audit = _require_exactly_one(audit_by_cell.get((sheet, raw_cell), []), f"formula audit {sheet}!{raw_cell}")
                if raw_formula != anchor["raw_G_formula"] or not _same_number(raw_cached, anchor["raw_G_cached"]):
                    raise ValueError(f"NEB inventory/workbook mismatch at {sheet}!{raw_cell}: formula/cache differ")
                if raw_audit["formula"] != raw_formula or not _same_number(raw_audit["cached_value"], raw_cached):
                    raise ValueError(f"formula audit/workbook mismatch at {sheet}!{raw_cell}")
                if int(raw_audit["row"]) != source_row:
                    raise ValueError(f"formula-audit row does not match source coordinate {sheet}!{raw_cell}")

                e0_formula, e0_cached = cell_values(sheet, anchor["E0_cell"])
                corr_formula, corr_cached = cell_values(sheet, anchor["Gibbs_correction_cell"])
                if not _same_number(e0_cached, anchor["E0"]):
                    raise ValueError(f"E0 cache differs at {sheet}!{anchor['E0_cell']}")
                if corr_formula != anchor["Gibbs_correction_formula_or_value"] or not _same_number(corr_cached, anchor["Gibbs_correction_cached"]):
                    raise ValueError(f"Gibbs correction differs at {sheet}!{anchor['Gibbs_correction_cell']}")

                mapped_formula_actual, mapped_cache_actual = cell_values(sheet, mapped_cell)
                mapped_audit = _require_exactly_one(audit_by_cell.get((sheet, mapped_cell), []), f"mapped formula audit {sheet}!{mapped_cell}")
                if mapped_audit["formula"] != mapped_formula_actual or not _same_number(mapped_audit["cached_value"], mapped_cache_actual):
                    raise ValueError(f"mapped-cell formula audit mismatch at {sheet}!{mapped_cell}")
                if mapped_expression.startswith("="):
                    if mapped_formula_actual != mapped_expression or not _same_number(mapped_cache_actual, mapped_energy):
                        raise ValueError(f"mapped-cell formula/value mismatch at {sheet}!{mapped_cell}")
                elif "spectator removed" in mapped_expression and "$F$60" in mapped_expression:
                    gas_formula, gas_cached = cell_values(sheet, "F60")
                    gas_manifest = json.loads((V1_DATA / "energy_mapping_manifest.json").read_text(encoding="utf-8"))["gas_reference_cached_values"]["C6H12O_g"]
                    if not _same_number(gas_cached, gas_manifest):
                        raise ValueError(f"C6H12O_g F60 cache differs from frozen mapping manifest: {gas_cached}")
                    normalized = Decimal(raw_cached) - Decimal(gas_cached)
                    if not _same_number(str(normalized), mapped_energy):
                        raise ValueError(f"Ti-Mo NEB2_FS normalization mismatch: {normalized} != {mapped_energy}")
                    mapped_expression = f"{mapped_expression}; raw-cell-formula={mapped_formula_actual}; F60-formula={gas_formula}; F60-cache={gas_cached}"
                else:
                    raise ValueError(f"unsupported adjusted NEB expression for {key}: {mapped_expression}")

                dft_value = dft_by_catalyst[catalyst][dft_field]
                if not _same_number(mapped_energy, dft_value):
                    raise ValueError(f"mapped energy differs from DFT reference for {key}: {mapped_energy} != {dft_value}")
                mapped_paths.setdefault((catalyst, path), {})[step_state] = Decimal(mapped_energy)
                generic_role = state["state_role"] if state else ""
                generic_label = state["state_label"] if state else ""
                if state is None:
                    role_status = "NO_GENERIC_INVENTORY_ROW_RAW_CELL_VERIFIED"
                elif generic_role != role or generic_label != anchor["state_label"]:
                    role_status = "ROLE_LABEL_CONFLICT_SOURCE_ROW_RETAINED"
                else:
                    role_status = "MATCH_SOURCE_ROW_AND_LABEL"
                next_formula, next_cache = cell_values(sheet, anchor["next_row_G_cell"])
                cr = {
                    "catalyst": catalyst, "source_workbook": workbook, "source_sheet": sheet,
                    "source_row": str(source_row), "NEB_role": role, "path": path, "state": step_state,
                    "raw_state_index_in_workbook": anchor["raw_state_index_in_workbook"],
                    "state_label_from_NEB_inventory": anchor["state_label"],
                    "generic_source_row_match_count": "1" if state else "0",
                    "generic_state_role": generic_role, "generic_state_label": generic_label,
                    "role_crosswalk_status": role_status,
                    "E0_cell": anchor["E0_cell"], "E0_formula_or_value": e0_formula, "E0_cached_eV": e0_cached,
                    "Gibbs_correction_cell": anchor["Gibbs_correction_cell"],
                    "Gibbs_correction_formula": corr_formula, "Gibbs_correction_cached_eV": corr_cached,
                    "raw_G_cell": raw_cell, "raw_G_formula": raw_formula, "raw_G_cached": raw_cached,
                    "formula_audit_match_count": "1", "next_row_G_cell": anchor["next_row_G_cell"],
                    "next_row_G_formula": next_formula, "next_row_G_cached": next_cache,
                    "mapped_neb_cell": mapped_cell, "mapped_neb_formula_or_adjustment": mapped_expression,
                    "mapped_neb_local_energy_eV": mapped_energy, "DFT_reference_value": dft_value,
                    "mapped_neb_energy_status": path_record["source_energy_crosswalk_status"],
                    "mapped_TS_minus_IS_eV": "", "mapped_TS_minus_FS_eV": "", "mapped_FS_minus_IS_eV": "",
                    "temperature_K": anchor["temperature_K"], "source_workbook_sha256": workbook_hash,
                    "DFT_TS_relative_sha256": dft_hash,
                }
                crosswalk.append(cr)
                anchor_crosswalk[(catalyst, role)] = cr
                label = anchor["state_label"]
                raw_state_index = anchor["raw_state_index_in_workbook"]
                source_path = anchor["calculation_paths"]
                temperature = anchor["temperature_K"]
                extra_cell = state.get("total_or_relative_cell", "") if state else ""
                generic_g_cell = state.get("G_cell", "") if state else ""
                generic_extra_cell = extra_cell
                mapping_status = path_record["source_energy_crosswalk_status"] + "|" + role_status
            else:
                role, label = state["state_role"], state["state_label"]
                raw_cell = state["G_cell"]
                raw_formula, raw_cached = cell_values(sheet, raw_cell)
                e0_formula, e0_cached = cell_values(sheet, state["E0_cell"])
                corr_formula, corr_cached = cell_values(sheet, state["Gibbs_correction_cell"])
                extra_cell = state["total_or_relative_cell"]
                generic_g_cell, generic_extra_cell = raw_cell, extra_cell
                raw_state_index, source_path = state["state_index"], state["calculation_path"]
                temperature = state["temperature_K"]
                mapping_status = "NOT_A_NEB_ANCHOR"
                mapped_cell = mapped_expression = mapped_energy = ""
                role_status = "GENERIC_ONLY_NO_CURATED_NEB_ROW"

            extra_formula, extra_cache = cell_values(sheet, extra_cell)
            term_formulas = [extra_formula]
            if anchor:
                term_formulas.extend((next_formula, mapped_expression))
            gas_terms = _gas_terms(*term_formulas)
            global_rows.append({
                "source_type": "NEB_ANCHOR" if anchor else "GENERIC_ONLY",
                "catalyst": catalyst, "source_workbook": workbook, "source_sheet": sheet,
                "source_row": str(source_row), "state_name": f"{role} | {label}", "state_role": role,
                "workbook_state_role": state["state_role"] if state else "",
                "workbook_state_label": state["state_label"] if state else "",
                "curated_neb_role": anchor["NEB_role"] if anchor else "",
                "role_crosswalk_status": role_status,
                "raw_state_index_in_workbook": raw_state_index, "source_file": str(SOURCE_WORKBOOK),
                "source_cell": raw_cell, "generic_inventory_G_cell": generic_g_cell,
                "generic_inventory_extra_cell": generic_extra_cell,
                "source_linux_calculation_path": source_path,
                "E0_cell": anchor["E0_cell"] if anchor else state["E0_cell"],
                "E0_formula_or_value": e0_formula, "E0_cached_eV": e0_cached,
                "temperature_K_as_recorded": temperature,
                "temperature_status": "RECORDED_333.15_K" if temperature == "333.15" else "MISSING_IN_WORKBOOK",
                "Gibbs_correction_cell": anchor["Gibbs_correction_cell"] if anchor else state["Gibbs_correction_cell"],
                "Gibbs_correction_formula": corr_formula, "Gibbs_correction_cached_eV": corr_cached,
                "original_G_formula": raw_formula, "original_G": raw_cached,
                "workbook_extra_cell": extra_cell, "workbook_extra_formula": extra_formula,
                "workbook_extra_cached_eV": extra_cache,
                "gas_reference_terms_named_in_extra_formula": " | ".join(gas_terms),
                "spectator_correction": "",
                "spectator_correction_status": "FORMULA_RECORDED_NOT_REAPPLIED" if gas_terms else "NO_EXPLICIT_GAS_TERM_IN_EXTRA_CELL_OR_UNRESOLVED",
                "mapped_neb_cell": mapped_cell, "mapped_neb_formula_or_adjustment": mapped_expression,
                "mapped_neb_local_energy_eV": mapped_energy, "mapped_neb_crosswalk_status": mapping_status,
                "slab_reference": "", "gas_reference": " | ".join(gas_terms),
                "reference_formula": REFERENCE_FORMULA, "calculated_formation_energy": "",
                "reference_status": "UNRESOLVED_NO_VERIFIED_CLEAN_SLAB_AND_ATOM_MAPPING",
                "missing_evidence": "clean-slab energy and structure; atom-resolved state composition; NEB block compatibility; frequency/thermal metadata where absent",
                "source_workbook_sha256": workbook_hash,
            })
    finally:
        formula_wb.close()
        cached_wb.close()

    expected_delta_fields = {
        "NEB1": ("dG_TS1_minus_IS1", "dG_TS1_minus_FS1"),
        "NEB2": ("dG_TS2_minus_IS2", "dG_TS2_minus_FS2"),
    }
    for (catalyst, path), values in mapped_paths.items():
        if set(values) != {"IS", "TS", "FS"}:
            raise ValueError(f"incomplete NEB path states for {catalyst}/{path}: {sorted(values)}")
        ts_is_field, ts_fs_field = expected_delta_fields[path]
        ts_minus_is, ts_minus_fs = values["TS"] - values["IS"], values["TS"] - values["FS"]
        dft_row = dft_by_catalyst[catalyst]
        if abs(ts_minus_is - Decimal(dft_row[ts_is_field])) > Decimal("1e-9"):
            raise ValueError(f"signed {path} TS-IS mismatch for {catalyst}")
        if abs(ts_minus_fs - Decimal(dft_row[ts_fs_field])) > Decimal("1e-9"):
            raise ValueError(f"signed {path} TS-FS mismatch for {catalyst}")
        anchor_crosswalk[(catalyst, f"{path}_TS")]["mapped_TS_minus_IS_eV"] = str(ts_minus_is)
        anchor_crosswalk[(catalyst, f"{path}_TS")]["mapped_TS_minus_FS_eV"] = str(ts_minus_fs)
        anchor_crosswalk[(catalyst, f"{path}_FS")]["mapped_FS_minus_IS_eV"] = str(values["FS"] - values["IS"])

    if len(mapped_paths) != 28 or len(crosswalk) != 84:
        raise ValueError(f"expected 28 paths and 84 NEB anchors; observed {len(mapped_paths)} and {len(crosswalk)}")
    matched_generic = sum(row["generic_source_row_match_count"] == "1" for row in crosswalk)
    conflicts = sum(row["role_crosswalk_status"] == "ROLE_LABEL_CONFLICT_SOURCE_ROW_RETAINED" for row in crosswalk)
    meta = {
        "source_sha256": workbook_hash, "expected_source_sha256": hashes[workbook_key],
        "dft_source_sha256": dft_hash, "expected_dft_source_sha256": hashes[dft_key],
        "state_rows": len(global_rows), "generic_inventory_rows": len(generic),
        "curated_neb_rows": len(neb), "crosswalk_rows": len(crosswalk),
        "generic_source_row_matches": matched_generic,
        "generic_source_row_missing": len(crosswalk) - matched_generic,
        "role_label_conflicts": conflicts, "neb_mapping_rows": len(path_by_key),
        "validated_paths": len(mapped_paths),
    }
    return global_rows, crosswalk, meta
