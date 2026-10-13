"""Replay Ti-Re NEB-block algebra and CatMAP static barrier semantics only."""

from __future__ import annotations

import csv
import json
import sys
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


V2_ROOT = Path(__file__).resolve().parents[1]
ROOT = V2_ROOT.parent
V1_DATA = ROOT / "data"
V1_MODEL = ROOT / "models" / "v1_serial"
DATA = V2_ROOT / "data"
REPORTS = V2_ROOT / "reports"
OUT_XLSX = DATA / "V2_TIRE_GLOBAL_REFERENCE_AUDIT.xlsx"
OUT_REPORT = REPORTS / "V2_TIRE_ENERGY_CYCLE.md"
sys.path.insert(0, str(V1_MODEL))


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def _d(value: str) -> Decimal:
    return Decimal(value)


def tire_rows() -> dict[str, dict[str, str]]:
    rows = {row["path"]: row for row in read_csv(V1_DATA / "neb_to_catmap_energy_mapping.csv")
            if row["catalyst"] == "tire"}
    if set(rows) != {"NEB1", "NEB2"}:
        raise ValueError("Ti-Re must have exactly one NEB1 and one NEB2 mapping")
    return rows


def replay() -> dict:
    from catmap.model import ReactionModel

    mappings = tire_rows()
    source_states = {(row["catalyst"], row["NEB_role"]): row for row in read_csv(V1_DATA / "original_neb_state_inventory.csv")}
    model = ReactionModel()
    expected = {
        "NEB1": ("Rpair_s", "TS1_s", "EpoxideHHO_s", Decimal("0.444846"), Decimal("2.372278"), Decimal("-1.927432")),
        "NEB2": ("HHO_s", "TS2_s", "OHHa_s", Decimal("0.217873"), Decimal("1.030708"), Decimal("-0.812835")),
    }
    results = []
    for path, (is_name, ts_name, fs_name, exp_fwd, exp_rev, exp_delta) in expected.items():
        row = mappings[path]
        g_is = _d(row["G_IS_workbook_or_normalized"])
        g_ts = _d(row["G_TS_workbook_or_normalized"])
        g_fs = _d(row["G_FS_workbook_or_normalized"])
        raw_fwd, raw_rev, raw_delta = g_ts - g_is, g_ts - g_fs, g_fs - g_is
        if (raw_fwd, raw_rev, raw_delta) != (exp_fwd, exp_rev, exp_delta):
            raise AssertionError(f"{path} signed differences changed: {(raw_fwd, raw_rev, raw_delta)}")
        energy = {is_name: float(g_is), ts_name: float(g_ts), fs_name: float(g_fs)}
        catmap_fwd = model.get_rxn_energy([[is_name], [ts_name], [fs_name]], energy)
        catmap_rev = model.get_rxn_energy([[fs_name], [ts_name], [is_name]], energy)
        if abs(catmap_fwd[1] - float(exp_fwd)) > 1e-12 or abs(catmap_rev[1] - float(exp_rev)) > 1e-12:
            raise AssertionError(f"CatMAP 0.3.1 static barrier replay failed for {path}")
        results.append({
            "path": path, "is_name": is_name, "ts_name": ts_name, "fs_name": fs_name,
            "row": row, "G_IS": g_is, "G_TS": g_ts, "G_FS": g_fs,
            "raw_TS_minus_IS": raw_fwd, "raw_TS_minus_FS": raw_rev, "raw_FS_minus_IS": raw_delta,
            "catmap_Ea_forward": catmap_fwd[1], "catmap_Ea_reverse": catmap_rev[1],
            "temperature_by_state": {
                is_name: source_states[("tire", f"{path}_IS")]["temperature_K"],
                ts_name: source_states[("tire", f"{path}_TS")]["temperature_K"],
                fs_name: source_states[("tire", f"{path}_FS")]["temperature_K"],
            },
            "status": "LOCAL_BLOCK_ALGEBRA_VERIFIED; GLOBAL_FORMATION_REFERENCE_UNRESOLVED",
        })
    return {"results": results, "catmap_version": __import__("catmap").__version__}


def _style(sheet) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="24476B")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for column in sheet.columns:
        values = [len(str(cell.value or "")) for cell in list(column)[:40]]
        width = min(48, max(12, max(values, default=10) + 2))
        sheet.column_dimensions[get_column_letter(column[0].column)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def write_xlsx(result: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "NEB local replay"
    headers = ["NEB", "State", "Excel source cell/formula", "G source eV", "T recorded K", "Local relative G eV",
               "Raw signed difference", "CatMAP effective Ea eV", "Evidence scope"]
    ws.append(headers)
    for item in result["results"]:
        row = item["row"]
        for state, key_cell, key_formula, g, local, raw, barrier in [
            (item["is_name"], "IS_source_cell", "IS_source_expression", item["G_IS"], Decimal(0), "", ""),
            (item["ts_name"], "TS_source_cell", "TS_source_expression", item["G_TS"], item["G_TS"]-item["G_IS"], item["raw_TS_minus_IS"], item["catmap_Ea_forward"]),
            (item["fs_name"], "FS_source_cell", "FS_source_expression_or_adjustment", item["G_FS"], item["G_FS"]-item["G_IS"], item["raw_FS_minus_IS"], ""),
        ]:
            ws.append([item["path"], state, f"{row[key_cell]} {row[key_formula]}", float(g),
                       item["temperature_by_state"][state], float(local), str(raw), barrier, item["status"]])
    _style(ws)
    cross = wb.create_sheet("Cross-path connection")
    cross.append(["Quantity", "Source/value", "Formula / interpretation", "Evidence status"])
    neb1 = result["results"][0]
    neb2 = result["results"][1]
    manifest = json.loads((V1_DATA / "energy_mapping_manifest.json").read_text(encoding="utf-8"))
    gas = Decimal(str(manifest["gas_reference_cached_values"]["C6H12O_g"]))
    endpoint = neb2["G_IS"] + gas
    delta = endpoint - neb1["G_FS"]
    cross.append(["NEB1 FS mapped G", float(neb1["G_FS"]), neb1["row"]["FS_source_expression_or_adjustment"], "NEB1 local mapping; common slab unverified"])
    cross.append(["NEB2 IS mapped G", float(neb2["G_IS"]), neb2["row"]["IS_source_expression"], "NEB2 local mapping; common slab unverified"])
    cross.append(["C6H12O gas G", float(gas), "F60; included once in conditional endpoint", "source cached gas reference"])
    cross.append(["post-desorption endpoint", float(endpoint), "G(NEB2 IS)+G(C6H12O_g)", "CONDITIONAL"])
    cross.append(["conditional ΔG", float(delta), "G(NEB2 IS)+G(C6H12O_g)-G(NEB1 FS)", "UNRESOLVED physically; not an activation barrier"])
    cross.append(["Global formation energy", None, "requires G_clean_slab and atom-resolved references", "UNRESOLVED; blank by design"])
    _style(cross)
    wb.save(OUT_XLSX)
    wb.close()


def write_report(result: dict) -> None:
    lines = [
        "# Ti–Re NEB局部能差、共同参考与能量循环审计", "",
        "## NEB1与NEB2内部差值复现", "",
        "| 路径 | IS/TS/FS来源单元格 | TS−IS | TS−FS | FS−IS | CatMAP有效正/逆Ea | 温度字段 | 局部数据状态 |",
        "|---|---|---:|---:|---:|---:|---|---|",
    ]
    for item in result["results"]:
        row = item["row"]
        cells = f"{row['IS_source_cell']}/{row['TS_source_cell']}/{row['FS_source_cell']}"
        temp_values = item["temperature_by_state"]
        temp = "/".join(temp_values[state] or "未知" for state in (item["is_name"], item["ts_name"], item["fs_name"]))
        lines.append(f"| {item['path']} | {cells} | {item['raw_TS_minus_IS']} | {item['raw_TS_minus_FS']} | {item['raw_FS_minus_IS']} | {item['catmap_Ea_forward']:.6f}/{item['catmap_Ea_reverse']:.6f} eV | {temp} | {item['status']} |")
    neb1, neb2 = result["results"]
    manifest = json.loads((V1_DATA / "energy_mapping_manifest.json").read_text(encoding="utf-8"))
    gas = Decimal(str(manifest["gas_reference_cached_values"]["C6H12O_g"]))
    endpoint = neb2["G_IS"] + gas
    delta = endpoint - neb1["G_FS"]
    lines.extend([
        "", "## Ti–Re NEB1/NEB2全局参考判定", "",
        "NEB1三状态相对能差精确重放为0.444846、2.372278、−1.927432 eV；NEB2为0.217873、1.030708、−0.812835 eV。CatMAP 0.3.1的`get_rxn_energy`静态重算得到相同有效正、逆势垒。此验证仅针对每条NEB内部已记录的相对能量。",
        "", "NEB1与NEB2局部G绝对值相差约100 eV，不能直接拼接为统一formation_energy。其数值差异本身也不能证明slab差异；共同slab、超胞和组成均缺结构证据。故全局形成能状态为**UNRESOLVED**。", "",
        f"环氧己烷脱附条件端点公式为`G(NEB2_IS)+G(C6H12O_g)-G(NEB1_FS)`，代入缓存值得到{delta} eV。它只在NEB1/NEB2共享slab、表面原子集合及温度修正时才是可比反应ΔG；当前只标为CONDITIONAL/UNRESOLVED，不是势垒。", "",
        f"工作簿气相参考`C6H12O_g`为F60={gas} eV；上式产品端只计入一次。Ti–Mo单独的`AR21=AR20+F60`及NEB2 FS减F60属于其路径局部映射规则，不能推广到Ti–Re，也不能未经构型核验再次应用。", "",
        "本审计只做十进制代数与CatMAP 0.3.1静态势垒函数调用；未运行速率、覆盖度、稳态或TOF。手工审阅表见`../data/V2_TIRE_GLOBAL_REFERENCE_AUDIT.xlsx`。",
    ])
    OUT_REPORT.write_text("\n".join(lines) + "\n", encoding="utf-8")


def build() -> dict:
    DATA.mkdir(parents=True, exist_ok=True)
    REPORTS.mkdir(parents=True, exist_ok=True)
    result = replay()
    write_xlsx(result)
    write_report(result)
    return {"catmap_version": result["catmap_version"], "paths_replayed": len(result["results"]),
            "global_reference_status": "UNRESOLVED", "tof_calculated": False}


if __name__ == "__main__":
    print(json.dumps(build(), ensure_ascii=False, indent=2))
