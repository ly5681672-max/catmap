"""Create an auditable source-to-formation-energy table without inventing offsets."""

from __future__ import annotations

import csv
import hashlib
import json
import re
from pathlib import Path

from openpyxl import load_workbook


V2_ROOT = Path(__file__).resolve().parents[1]
ROOT = V2_ROOT.parent
V1_DATA = ROOT / "data"
DATA = V2_ROOT / "data"
REPORTS = V2_ROOT / "reports"
SOURCE_WORKBOOK = Path(r"H:\catmap\catmap\lyq_alkene_epoxidation\己烯环氧化.xlsx")
GLOBAL_CSV = DATA / "V2_GLOBAL_GIBBS_REFERENCE.csv"
FORMATION_REPORT = REPORTS / "V2_FORMATION_ENERGY_DERIVATION.md"
REFERENCE_FORMULA = "G_full(X*) - G_clean_slab - sum_i(n_i * mu_i_ref)"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest().upper()


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def _normalize(value):
    return "" if value is None else value


def _gas_terms(formula: str) -> list[str]:
    terms = []
    if "F59" in formula:
        terms.append("H2O_g:F59")
    if "F60" in formula:
        terms.append("C6H12O_g:F60")
    return terms


def _mapping_for_state(row: dict[str, str], mapping: dict[tuple[str, str], dict[str, str]]) -> dict[str, str]:
    role = row["NEB_role"]
    match = re.fullmatch(r"(NEB1|NEB2)_(IS|TS|FS)", role)
    if not match:
        return {"mapped_cell": "", "mapped_formula": "", "mapped_energy": "", "mapping_status": "NOT_A_NEB_ANCHOR"}
    path, state = match.groups()
    record = mapping.get((row["catalyst"], path))
    if not record:
        return {"mapped_cell": "", "mapped_formula": "", "mapped_energy": "", "mapping_status": "MISSING_MAPPING"}
    if state == "IS":
        fields = ("IS_source_cell", "IS_source_expression", "G_IS_workbook_or_normalized")
    elif state == "TS":
        fields = ("TS_source_cell", "TS_source_expression", "G_TS_workbook_or_normalized")
    else:
        fields = ("FS_source_cell", "FS_source_expression_or_adjustment", "G_FS_workbook_or_normalized")
    return {
        "mapped_cell": record[fields[0]], "mapped_formula": record[fields[1]],
        "mapped_energy": record[fields[2]],
        "mapping_status": record["source_energy_crosswalk_status"],
    }


def build_rows() -> tuple[list[dict[str, str]], dict]:
    inventory = read_csv(V1_DATA / "original_gibbs_state_inventory.csv")
    neb_inventory = read_csv(V1_DATA / "original_neb_state_inventory.csv")
    mapping_rows = read_csv(V1_DATA / "neb_to_catmap_energy_mapping.csv")
    mapping = {(row["catalyst"], row["path"]): row for row in mapping_rows}
    source_manifest = json.loads((V1_DATA / "source_hashes_before.json").read_text(encoding="utf-8"))
    source_hashes = source_manifest["sha256"]
    source_key = str(SOURCE_WORKBOOK)
    if source_key not in source_hashes:
        raise ValueError(f"Original workbook path is not listed in the frozen source manifest: {SOURCE_WORKBOOK}")
    actual_hash = sha256(SOURCE_WORKBOOK)
    expected_hash = source_hashes[source_key]
    if actual_hash != expected_hash:
        raise ValueError(f"Source workbook hash mismatch: expected {expected_hash}, observed {actual_hash}")

    formula_wb = load_workbook(SOURCE_WORKBOOK, read_only=True, data_only=False)
    cached_wb = load_workbook(SOURCE_WORKBOOK, read_only=True, data_only=True)
    neb_lookup = {(row["catalyst"], row["NEB_role"]): row for row in neb_inventory}
    output: list[dict[str, str]] = []
    for row in inventory:
        sheet_name = row["worksheet"]
        formula_sheet = formula_wb[sheet_name]
        cached_sheet = cached_wb[sheet_name]

        def values(cell_ref: str) -> tuple[str, str]:
            if not cell_ref:
                return "", ""
            return str(_normalize(formula_sheet[cell_ref].value)), str(_normalize(cached_sheet[cell_ref].value))

        e0_formula, e0_cached = values(row["E0_cell"])
        correction_formula, correction_cached = values(row["Gibbs_correction_cell"])
        g_formula, g_cached = values(row["G_cell"])
        extra_formula, extra_cached = values(row["total_or_relative_cell"])
        neb = neb_lookup.get((row["catalyst"], row["state_role"]), {})
        mapped = _mapping_for_state(neb, mapping) if neb else {
            "mapped_cell": "", "mapped_formula": "", "mapped_energy": "", "mapping_status": "NOT_A_NEB_ANCHOR"
        }
        gas_terms = _gas_terms(extra_formula)
        output.append({
            "catalyst": row["catalyst"],
            "state_name": f"{row['state_role']} | {row['state_label']}",
            "state_role": row["state_role"],
            "workbook_state_label": row["state_label"],
            "source_file": str(SOURCE_WORKBOOK),
            "source_sheet": sheet_name,
            "source_cell": row["G_cell"],
            "source_linux_calculation_path": row["calculation_path"],
            "E0_cell": row["E0_cell"],
            "E0_formula_or_value": e0_formula,
            "E0_cached_eV": e0_cached,
            "temperature_K_as_recorded": row["temperature_K"],
            "temperature_status": "RECORDED_333.15_K" if row["temperature_K"] == "333.15" else "MISSING_IN_WORKBOOK",
            "Gibbs_correction_cell": row["Gibbs_correction_cell"],
            "Gibbs_correction_formula": correction_formula,
            "Gibbs_correction_cached_eV": correction_cached,
            "original_G_formula": g_formula,
            "original_G": g_cached,
            "workbook_extra_cell": row["total_or_relative_cell"],
            "workbook_extra_formula": extra_formula,
            "workbook_extra_cached_eV": extra_cached,
            "gas_reference_terms_named_in_extra_formula": " | ".join(gas_terms),
            "spectator_correction": "",
            "spectator_correction_status": "FORMULA_RECORDED_NOT_REAPPLIED" if gas_terms else "NO_EXPLICIT_GAS_TERM_IN_EXTRA_CELL_OR_UNRESOLVED",
            "mapped_neb_cell": mapped["mapped_cell"],
            "mapped_neb_formula_or_adjustment": mapped["mapped_formula"],
            "mapped_neb_local_energy_eV": mapped["mapped_energy"],
            "mapped_neb_crosswalk_status": mapped["mapping_status"],
            "slab_reference": "",
            "gas_reference": " | ".join(gas_terms),
            "reference_formula": REFERENCE_FORMULA,
            "calculated_formation_energy": "",
            "reference_status": "UNRESOLVED_NO_VERIFIED_CLEAN_SLAB_AND_ATOM_MAPPING",
            "missing_evidence": "clean-slab energy and structure; atom-resolved state composition; NEB block compatibility; frequency/thermal metadata where absent",
            "source_workbook_sha256": actual_hash,
        })
    formula_wb.close()
    cached_wb.close()
    return output, {"source_sha256": actual_hash, "expected_source_sha256": expected_hash,
                    "state_rows": len(output), "neb_mapping_rows": len(mapping_rows)}


def build() -> dict:
    rows, meta = build_rows()
    DATA.mkdir(parents=True, exist_ok=True)
    REPORTS.mkdir(parents=True, exist_ok=True)
    with GLOBAL_CSV.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)

    timo1 = next(row for row in rows if row["catalyst"] == "timo" and row["state_role"] == "NEB1_FS")
    timo2 = next(row for row in rows if row["catalyst"] == "timo" and row["state_role"] == "NEB2_FS")
    report = f"""# V2全局Gibbs形成能推导与证据边界

## 数据来源与分层

逐状态读取主工作簿`{SOURCE_WORKBOOK}`的公式视图与缓存值视图，并与V1已冻结的状态清单、NEB映射和输入哈希交叉检查。工作簿SHA256为`{meta['source_sha256']}`，与冻结值匹配。输出共{meta['state_rows']}条工作簿状态；NEB映射表包含{meta['neb_mapping_rows']}条路径记录（14个位点×2条NEB）。

本报告分别保留：(1) Excel状态单元格的E0；(2) Gibbs修正单元格中的原始值/公式及缓存；(3) `G=E0+correction`单元格的原始公式及缓存；(4) 旁列“总G/相对量”的原始公式和缓存；(5) V1 NEB逐路径使用的映射值/特殊调整。上述字段不互相替代，也不自动视为CatMAP全局formation_energy。

## 当前形成能定义

\[G_{{form}}(X^*)=G_{{full}}(X^*)-G_{{clean\ slab}}-\sum_i n_i\mu_i^{{ref}}.\]

当前没有可追溯的clean-slab结构/能量及逐状态原子映射；因此每行`calculated_formation_energy`保持空白。空白不等于0，也不以NEB局部IS为全局零点。Excel状态标签或研究者确认可用作化学状态解释，但不能填补Slab参考或坐标缺口。

## Ti–Mo旁观气相项专项核查

- NEB1 FS的V1相对能映射使用`AR21 = AR20 + F60`，即保留工作簿显式加入的`C6H12O_g`气相项；`AR20`原始G仍单独记录。不能再把这一映射值误当作未调整的原始态G。
- NEB2 FS使用`AR32 - F60`移除Excel比较项中的一份`C6H12O_g`，以复现CSV给定的NEB2局部相对差。
- 这些是NEB路径局部交叉核验规则；在两态结构和原子清单不可见时，不据此推出全局脱附ΔG，也不把`F60`再次自动加/减到其他状态。

V1来源公式：NEB1 FS映射公式=`{timo1['mapped_neb_formula_or_adjustment']}`；NEB2 FS映射公式=`{timo2['mapped_neb_formula_or_adjustment']}`。两条映射保留在`V2_GLOBAL_GIBBS_REFERENCE.csv`的独立字段中。

## 证据状态

- Ti–Re NEB1 IS/TS/FS工作簿记录温度为333.15 K；NEB2 IS/TS温度单元格为空，NEB2 FS为333.15 K。NEB2 IS/TS的修正公式及缓存虽存在，但不能据此填补温度或重建热化学来源。
- 对同一NEB路径内部的原始有符号能差，可以复现工作簿和CSV的局部关系；这不证明两条路径共享slab参考。
- 所有状态的全局formation_energy均未计算，状态映射和缺失证据详见CSV及结构审计。

结论：在当前可用资料下，统一参考转换**不能数值完成**；本项目只交付可审计原始数据层、局部映射层和缺失项，不引入人为平移。
"""
    FORMATION_REPORT.write_text(report, encoding="utf-8")
    return meta


if __name__ == "__main__":
    print(json.dumps(build(), ensure_ascii=False, indent=2))
