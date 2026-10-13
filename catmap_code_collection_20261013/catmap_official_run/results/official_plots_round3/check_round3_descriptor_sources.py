"""Audit real C/O/H adsorbate sources and CatMAP 0.3.1 scaler eligibility.

This script does not run CatMAP kinetics or create/interpolate plots. It writes
the source-cell crosswalk and verifies whether the proposed atomic adsorbates
are present in the unchanged six-step reaction network.
"""

from __future__ import annotations

import csv
import os
from pathlib import Path

from openpyxl import load_workbook
from catmap import ReactionModel, __version__


OUT_DIR = Path(__file__).resolve().parent
RUN_DIR = OUT_DIR.parents[1]
WORKBOOK = Path(r"H:\Doctoral_research\己烯环氧化.xlsx")
MODEL_DIR = RUN_DIR / "results" / "official_plots_round2"
MODEL_FILE = MODEL_DIR / "alkene_epoxidation_descriptor_map.mkm"
CSV_OUT = OUT_DIR / "descriptor_source_crosswalk.csv"

SURFACES = {
    "TiFe": "tife", "TiMn": "timn", "TiHf": "tihf", "TiRe": "tire",
    "TiMo": "timo", "TiV": "tiv", "TiCo": "tico", "TiTi": "titi",
    "TiW": "tiw", "TiTa": "tita", "TiCr": "ticr", "单核Ti": "ti",
    "TiNi": "tini", "TiZr": "tizr",
}
STATE_LABELS = {"H_a_111": "H_a", "H_b_111": "H_b", "O_111": "O"}
PAIR_CHECKS = {
    "C-O": ("C_atom_s", "O_s"),
    "C-Ha": ("C_atom_s", "H_a_s"),
    "O-Ha": ("O_s", "H_a_s"),
}


def xcell(ws, row: int, col: int) -> str:
    return ws.cell(row=row, column=col).coordinate


def scaler_rejection(pair: tuple[str, str]) -> str:
    old_cwd = Path.cwd()
    try:
        os.chdir(MODEL_DIR)
        model = ReactionModel(setup_file=str(MODEL_FILE))
        if len(model.elementary_rxns) != 6:
            raise RuntimeError("The existing model did not parse as six reactions")
        model.descriptor_names = list(pair)
        try:
            model.scaler.parameterize()
        except Exception as exc:  # report official scaler's exact diagnostic
            return f"{type(exc).__name__}: {exc}"
        return "SUPPORTED"
    finally:
        os.chdir(old_cwd)


def main() -> None:
    if __version__ != "0.3.1":
        raise RuntimeError(f"Expected CatMAP 0.3.1; found {__version__}")
    if not WORKBOOK.is_file() or not MODEL_FILE.is_file():
        raise FileNotFoundError("Expected original workbook or round2 model is missing")

    formulas = load_workbook(WORKBOOK, data_only=False, read_only=True)
    cached = load_workbook(WORKBOOK, data_only=True, read_only=True)
    wf, wc = formulas["catmap数据集"], cached["catmap数据集"]
    rows: list[dict[str, str]] = []

    for col in range(1, wf.max_column + 1):
        catalyst = wf.cell(1, col).value
        surface = SURFACES.get(str(catalyst))
        if not surface:
            continue
        states: dict[str, tuple[int, int]] = {}
        for row in range(1, wf.max_row + 1):
            label = wf.cell(row, col).value
            if label in STATE_LABELS:
                states[STATE_LABELS[label]] = (row, row + 2)

        for requested, state in (("O", "O"), ("H", "H_a"), ("H_alternative", "H_b")):
            if state not in states:
                rows.append({
                    "surface_name": surface, "catalyst_header": str(catalyst),
                    "requested_descriptor": requested, "source_species": "",
                    "status": "SOURCE_STATE_NOT_FOUND",
                })
                continue
            label_row, data_row = states[state]
            e0_col, dg_col, g_col = col + 1, col + 4, col + 5
            rows.append({
                "surface_name": surface,
                "catalyst_header": str(catalyst),
                "requested_descriptor": requested,
                "source_species": f"{state}_111",
                "status": "SOURCE_AVAILABLE_BUT_NOT_IN_CURRENT_NETWORK",
                "state_label_cell": xcell(wf, label_row, col),
                "E0_cell": xcell(wf, data_row, e0_col),
                "E0_eV": str(wc.cell(data_row, e0_col).value),
                "deltaG_cell": xcell(wf, data_row, dg_col),
                "deltaG_formula": str(wf.cell(data_row, dg_col).value),
                "deltaG_cached_eV": str(wc.cell(data_row, dg_col).value),
                "G_cell": xcell(wf, data_row, g_col),
                "G_formula": str(wf.cell(data_row, g_col).value),
                "G_cached_eV": str(wc.cell(data_row, g_col).value),
                "temperature_cell": xcell(wf, data_row, col + 2),
                "temperature_value": str(wc.cell(data_row, col + 2).value),
                "path_cell": xcell(wf, data_row, col + 3),
                "path_value": str(wc.cell(data_row, col + 3).value),
                "note": "Workbook Gibbs source only; not a CatMAP formation_energy.",
            })

        # The requested atomic carbon adsorbate must be explicit. C6H12_111 is
        # deliberately not treated as C*.
        c_labels = []
        for row in range(1, wf.max_row + 1):
            label = wf.cell(row, col).value
            if isinstance(label, str) and label.strip().lower() in {
                "c_111", "c*_111", "c_atom_111", "c_s"
            }:
                c_labels.append((row, label))
        rows.append({
            "surface_name": surface,
            "catalyst_header": str(catalyst),
            "requested_descriptor": "C",
            "source_species": ";".join(f"{label}@{xcell(wf,row,col)}" for row,label in c_labels),
            "status": "SOURCE_AVAILABLE" if c_labels else "NO_INDEPENDENT_C_ADSORBATE",
            "note": "C6H12_111 is present but intentionally excluded; it is not C*.",
        })

    formulas.close()
    cached.close()

    with CSV_OUT.open("w", newline="", encoding="utf-8-sig") as stream:
        fields = sorted({key for row in rows for key in row})
        writer = csv.DictWriter(stream, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)

    print(f"CatMAP {__version__}; source crosswalk rows={len(rows)}; {CSV_OUT}")
    for pair_name, pair in PAIR_CHECKS.items():
        print(f"{pair_name}: {scaler_rejection(pair)}")
    print("No kinetic solve or plotting was performed.")


if __name__ == "__main__":
    main()
