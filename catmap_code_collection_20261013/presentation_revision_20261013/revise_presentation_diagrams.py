from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor


BASE = Path(r"C:\Users\Administrator\Documents\codex探索\catmap_neb_coadsorption_model\presentation_revision_20261013")
SOURCE = BASE / "汇报2026_9_9_官网CatMAP框架修订版_20261013_修订版.pptx"
OUTPUT = BASE / "汇报2026_9_9_官网CatMAP框架修订版_20261013_图示简洁版.pptx"


def replace_text(shape, text, size_pt=None):
    tf = shape.text_frame
    ps = list(tf.paragraphs)
    p0 = ps[0]
    old = p0.runs[0] if p0.runs else None
    font = old.font if old else None
    name = font.name if font else None
    bold = font.bold if font else None
    italic = font.italic if font else None
    color = None
    if font:
        try:
            if font.color.type is not None and font.color.rgb is not None:
                color = font.color.rgb
        except (AttributeError, ValueError):
            pass
    p0.clear()
    for p in ps[1:]:
        p._element.getparent().remove(p._element)
    r = p0.add_run()
    r.text = text
    if name:
        r.font.name = name
    if bold is not None:
        r.font.bold = bold
    if italic is not None:
        r.font.italic = italic
    if color is not None:
        r.font.color.rgb = color
    if size_pt is not None:
        r.font.size = Pt(size_pt)
    tf.word_wrap = True


prs = Presentation(str(SOURCE))

# Slide 2 — retain the three panel layout, clear dense interior copy, and show the research path.
s = prs.slides[1]
replace_text(s.shapes[0], "从原始计算到可检验描述符：当前研究路径", 22)
replace_text(s.shapes[8], "原始计算", 17)
replace_text(s.shapes[9], "官方CatMAP", 17)
replace_text(s.shapes[10], "描述符检验", 17)
replace_text(s.shapes[13], "输入", 14)
replace_text(s.shapes[14], "求解", 14)
replace_text(s.shapes[15], "检验", 14)
to_remove = [s.shapes[idx] for idx in (11, 12, 16, 17, 18, 19, 20, 21, 22, 23, 24, 25, 26, 27, 28, 29, 30, 31, 32)]
for shape in to_remove:
    shape._element.getparent().remove(shape._element)

def add_panel_copy(slide, x, heading, detail, foot):
    box = slide.shapes.add_textbox(Inches(x), Inches(2.45), Inches(3.45), Inches(2.65))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.12)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.15)
    tf.margin_bottom = Inches(0.1)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, (value, size, bold, color) in enumerate([
        (heading, 22, True, RGBColor(36, 54, 86)),
        (detail, 17, True, RGBColor(50, 78, 110)),
        (foot, 13, False, RGBColor(70, 78, 86)),
    ]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.space_before = Pt(8 if i else 0)
        r = p.add_run()
        r.text = value
        r.font.name = "Microsoft YaHei"
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color

add_panel_copy(s, 0.85, "DFT + NEB", "Gᵢ  /  ΔG‡", "原始能量与来源可追溯")
add_panel_copy(s, 4.95, "CatMAP 0.3.1", "6步网络  →  θ_ss  →  TOF", "13个位点已复核；Ti–Zr未验证")
add_panel_copy(s, 9.05, "Ti–M电子结构", "?  →  ΔG‡_NEB1  →  TOF", "候选关系待检验，不预设因果")
for x in (4.53, 8.77):
    arrow = s.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x), Inches(3.45), Inches(0.22), Inches(0.55))
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = RGBColor(65, 105, 145)
    arrow.line.fill.background()

# Slide 5 — turn the dense prose into a short four-stage equation/flow.
s = prs.slides[4]
replace_text(s.shapes[2], "官方稳态主线：输入能量与条件 → 求解覆盖度稳态 → 输出TOF与速率图", 16)
replace_text(s.shapes[12], "TableParser\n读取物种、形成能与反应网络\nGibbs采用frozen模式", 13)
replace_text(s.shapes[17], "CatMAP速率\nr_j = r_j(θ; k_j, T, p)\n输入条件保持固定", 13)
replace_text(s.shapes[22], "稳态求解\nS · r(θ) = 0\n得到 θ_ss", 13)
replace_text(s.shapes[27], "官方映射\nθ_ss → TOF / production rate\n输出 descriptor maps", 13)
replace_text(s.shapes[28], "核心：S · r(θ_ss) = 0\n这不是本项目的时间积分轨迹。", 14)
replace_text(s.shapes[29], "CatMAP 0.3.1：ReactionModel → TableParser / ThermodynamicScaler → SteadyStateSolver → Mapper / VectorMap", 12)

# Slide 6 — retain the native CatMAP map; reduce the explanation to a compact legend.
s = prs.slides[5]
replace_text(s.shapes[2], "13个位点的CatMAP官方二维产物TOF图（探索性）", 15)
replace_text(s.shapes[4], "读图", 16)
replace_text(
    s.shapes[5],
    "横轴：TS1 formation energy (eV)\n纵轴：EpoxideHHO formation energy (eV)\n颜色：log10(TOF)\n\n13个位点含名称；Ti–Zr未验证，不绘为零。",
    15,
)

# Slide 7 — express evidence level and model caveat as short visual-card labels.
s = prs.slides[6]
replace_text(s.shapes[1], "CatMAP第一轮结果：可比较范围与模型边界", 23)
replace_text(s.shapes[2], "统一条件下的探索性结果，不等同于完整过渡态验证后的最终活性排序。", 15)
replace_text(s.shapes[6], "机理边界")
replace_text(s.shapes[7], "六步网络\nTS1 / TS2：原始NEB\nHHO*：分散共吸附态", 15)
replace_text(s.shapes[11], "位点状态")
replace_text(s.shapes[12], "13个位点\n稳态与通量已复核\nTi–Zr未验证", 15)
replace_text(s.shapes[16], "解释边界")
replace_text(s.shapes[17], "步骤1、3、5、6：端点近似\nΔG3–TOF关联受该假设影响\n不是普适因果证明", 14)
replace_text(s.shapes[18], "结论：模型内可比较；外推需补强关键势垒证据。", 15)

# Slide 8 — a compact four-step next-work diagram, avoiding paragraph blocks.
s = prs.slides[7]
replace_text(s.shapes[1], "后续研究：从模型响应到可检验机理", 23)
replace_text(s.shapes[2], "保持六步网络、原始DFT与333.15 K条件；先验证物理关系，再考虑拟合。", 15)
replace_text(s.shapes[6], "固定模型边界")
replace_text(s.shapes[7], "六步网络不变\nTS1 / TS2：显式NEB\n其他四步：端点近似", 15)
replace_text(s.shapes[11], "整理电子结构")
replace_text(s.shapes[12], "已有Ti–M特征\n↓\n与NEB1势垒 / TOF关联待检", 15)
replace_text(s.shapes[16], "补强势垒证据")
replace_text(s.shapes[17], "优先关注\n步骤1 / 3 / 5 / 6\n不填造TS能量", 15)
replace_text(s.shapes[21], "保留证据分级")
replace_text(s.shapes[22], "13个位点：探索性\nTi–Zr：未解决\n失败不记为零活性", 15)
replace_text(s.shapes[23], "成功标准：结论可追溯、可检验；不预设火山峰。", 15)

# Slide 9 — replace the generic thesis/database page with the actual literature-to-database workflow.
s = prs.slides[8]
old_shapes = {i: s.shapes[i] for i in (2, 3, 4, 5, 6, 10, 11)}
for sh in old_shapes.values():
    sh._element.getparent().remove(sh._element)
replace_text(s.shapes[0], "文献数据资源", 16)
replace_text(s.shapes[1], "从文献证据到结构化催化科学数据库", 23)

cards = [
    ("文献来源", "论文 / Supporting Information\n保留原文与出处", "EAF1F8"),
    ("识别与抽取", "正文、表格与实验条件\n定位到页 / 表 / 单元格", "E6F3F0"),
    ("证据校核", "核对数值、单位、对象与条件\n歧义项隔离待复核", "FFF1DE"),
    ("结构化利用", "反应 / 性能 / TOF记录\nscientific.db → 检索与分析", "EEEAF7"),
]
xs = [0.55, 3.65, 6.75, 9.85]
for idx, (title, body, fill) in enumerate(cards):
    card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(xs[idx]), Inches(2.45), Inches(2.78), Inches(2.2))
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor.from_string(fill)
    card.line.color.rgb = RGBColor(124, 145, 165)
    tf = card.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.12)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.2)
    tf.margin_bottom = Inches(0.12)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = title
    r.font.name = "Microsoft YaHei"
    r.font.size = Pt(19)
    r.font.bold = True
    r.font.color.rgb = RGBColor(36, 54, 86)
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(12)
    r2 = p2.add_run()
    r2.text = body
    r2.font.name = "Microsoft YaHei"
    r2.font.size = Pt(13)
    r2.font.color.rgb = RGBColor(55, 65, 75)
    if idx < 3:
        arrow = s.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(xs[idx] + 2.83), Inches(3.28), Inches(0.24), Inches(0.52))
        arrow.fill.solid()
        arrow.fill.fore_color.rgb = RGBColor(65, 105, 145)
        arrow.line.fill.background()

formula = s.shapes.add_textbox(Inches(0.8), Inches(5.12), Inches(11.75), Inches(0.72))
tf = formula.text_frame
tf.clear()
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "一条科学记录 = 数值 + 单位 + 实验条件 + 可追溯文献定位"
r.font.name = "Microsoft YaHei"
r.font.size = Pt(20)
r.font.bold = True
r.font.color.rgb = RGBColor(36, 54, 86)

prs.save(str(OUTPUT))
print(f"SAVED {OUTPUT}")
