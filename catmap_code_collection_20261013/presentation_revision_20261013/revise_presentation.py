from pathlib import Path
from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor


BASE = Path(r"C:\Users\Administrator\Documents\codex探索\catmap_neb_coadsorption_model\presentation_revision_20261013")
DECK = BASE / "汇报2026_9_9_官网CatMAP框架修订版_20261013_修订版.pptx"
SOURCE = BASE / "汇报2026_9_9_官网CatMAP框架修订版_20261013_原稿备份.pptx"


def replace_text(shape, text, size_pt=None):
    """Replace text while retaining the existing text-frame and first-run styling."""
    tf = shape.text_frame
    paragraphs = list(tf.paragraphs)
    p0 = paragraphs[0]
    old_run = p0.runs[0] if p0.runs else None
    old_font = old_run.font if old_run else None
    font_name = old_font.name if old_font else None
    font_bold = old_font.bold if old_font else None
    font_italic = old_font.italic if old_font else None
    font_color = None
    if old_font:
        try:
            if old_font.color.type is not None and old_font.color.rgb is not None:
                font_color = old_font.color.rgb
        except (AttributeError, ValueError):
            pass
    p0.clear()
    for p in paragraphs[1:]:
        p._element.getparent().remove(p._element)
    run = p0.add_run()
    run.text = text
    if font_name:
        run.font.name = font_name
    if font_bold is not None:
        run.font.bold = font_bold
    if font_italic is not None:
        run.font.italic = font_italic
    if font_color is not None:
        run.font.color.rgb = font_color
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    tf.word_wrap = True


prs = Presentation(str(SOURCE))

# Slide 2: remove the unfinished SISSO placeholder without implying a fitted model.
replace_text(
    prs.slides[1].shapes[28],
    "SISSO（后续）\n候选特征待检验\n当前无已验证回归式",
    13,
)

# Slide 3: replace the obsolete eight-step mechanism with the approved six-step network.
replace_text(prs.slides[2].shapes[0], "己烯环氧化：CatMAP六步串联机理（333.15 K）", 22)
replace_text(
    prs.slides[2].shapes[1],
    "* + H2O2_g + C6H12_g <-> Rpair*\n"
    "Rpair* <-> TS1* -> EpoxideHHO*\n"
    "EpoxideHHO* <-> HHO* + C6H12O_g\n"
    "HHO* <-> TS2* -> OHHa*\n"
    "OHHa* <-> H2O*\n"
    "H2O* <-> H2O_g + *\n\n"
    "总反应：H2O2_g + C6H12_g -> C6H12O_g + H2O_g\n"
    "TS1/TS2对应原始NEB；第1、3、5、6步采用CatMAP官方端点势垒近似。",
    17,
)
replace_text(prs.slides[2].shapes[4], "文献背景（非本项目反应网络）", 11)

# Slide 4: replace the obsolete eight-step screenshot with editable CatMAP expressions.
slide4 = prs.slides[3]
old_image = slide4.shapes[0]
slide4_title = slide4.shapes[1]
slide4_caption = slide4.shapes[6]
left, top, width, height = old_image.left, old_image.top, old_image.width, old_image.height
old_image._element.getparent().remove(old_image._element)
replace_text(slide4_title, "CatMAP 0.3.1：完整共吸附态的六步串联表示", 22)
code_box = slide4.shapes.add_textbox(left, top, width, height)
tf = code_box.text_frame
tf.clear()
tf.word_wrap = False
tf.margin_left = Inches(0.24)
tf.margin_right = Inches(0.12)
tf.margin_top = Inches(0.12)
tf.margin_bottom = Inches(0.06)
tf.vertical_anchor = MSO_ANCHOR.TOP
lines = [
    "rxn_expressions = [",
    "  '* + H2O2_g + C6H12_g <-> Rpair*',",
    "  'Rpair* <-> TS1* -> EpoxideHHO*',",
    "  'EpoxideHHO* <-> HHO* + C6H12O_g',",
    "  'HHO* <-> TS2* -> OHHa*',",
    "  'OHHa* <-> H2O*',",
    "  'H2O* <-> H2O_g + *',",
    "]",
]
for i, line in enumerate(lines):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.text = line
    p.space_after = Pt(4)
    for run in p.runs:
        run.font.name = "Courier New"
        run.font.size = Pt(18)
        run.font.color.rgb = RGBColor(36, 54, 86)
note = tf.add_paragraph()
note.text = "说明：Rpair、EpoxideHHO、HHO、OHHa均按完整共吸附态表示；HHO*不是分子态水。仅TS1/TS2为显式NEB过渡态。"
note.space_before = Pt(8)
for run in note.runs:
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(13)
    run.font.italic = True
    run.font.color.rgb = RGBColor(80, 80, 80)
replace_text(
    slide4_caption,
    "第1、3、5、6步采用CatMAP官方端点势垒近似；此版本属于探索性微观动力学模型。",
    13,
)

# Slide 5: align the overview with the actual frozen Gibbs input and official workflow.
replace_text(prs.slides[4].shapes[1], "CatMAP 0.3.1官方稳态求解与描述符映射流程", 22)
replace_text(
    prs.slides[4].shapes[12],
    "ReactionModel / TableParser 读取 .mkm 与 energies.txt；已输入的Excel Gibbs采用 frozen 模式，避免重复热修正。\n解决：科学能量如何进入官方模型。",
    12,
)
replace_text(
    prs.slides[4].shapes[17],
    "CatMAP根据覆盖度、气相输入和势垒构造各基元速率 r_j。\n解决：能量与势垒如何转化为覆盖度相关速率。",
    12,
)
replace_text(
    prs.slides[4].shapes[27],
    "Mapper / VectorMap 扫描描述符网格并输出官方 rate、coverage、production-rate maps。\n当前二维TOF图采用 TS1_s 与 EpoxideHHO_s formation energy。",
    12,
)
replace_text(
    prs.slides[4].shapes[29],
    "官方流程：ReactionModel → TableParser / ThermodynamicScaler → SteadyStateSolver → Mapper / VectorMap；具体结果仍须结合收敛与模型假设分级解释。",
    11,
)

# Slide 6: describe the existing official map accurately; its 13 catalyst labels remain intact.
replace_text(prs.slides[5].shapes[1], "CatMAP官方二维产物TOF图：13个位点的描述符空间分布", 22)
replace_text(
    prs.slides[5].shapes[2],
    "图示为当前模型的探索性响应，不代表已验证的普适火山关系或因果描述符。",
    13,
)
replace_text(prs.slides[5].shapes[4], "官方二维网格与真实催化剂点", 16)
replace_text(
    prs.slides[5].shapes[5],
    "颜色：C6H12O_g 的 log10(TOF)；色标约为 −11 至 −8。\n"
    "横轴：TS1_s formation energy (eV)。\n纵轴：EpoxideHHO_s formation energy (eV)。\n\n"
    "黑点及名称：13个完成稳态/通量复核的位点（单Ti及12种双核位点），已完整叠加。\n"
    "Ti–Zr未通过稳态验证，故不作为有效点绘入。\n\n"
    "此图展示固定333.15 K及既定储库下的CatMAP官方描述符扫描；步骤1、3、5、6使用端点势垒近似，结果应按探索性证据解释。",
    13,
)

# Slide 7: replace stale rejection language with the current evidence status and limitations.
replace_text(prs.slides[6].shapes[1], "第一轮CatMAP结果：13个位点已复核，Ti–Zr仍未解决", 23)
replace_text(
    prs.slides[6].shapes[2],
    "当前结论：可比较的是统一模型假设下的探索性TOF；不等同于完整过渡态验证后的最终活性排序。",
    14,
)
replace_text(prs.slides[6].shapes[7],
    "六步网络保持完整共吸附态。\n\nTS1、TS2来自原始NEB；HHO*表示分散吸附的Ha/Hb/O，不是H2O分子。\n\n未把网络外的C*、O*、Ha*或Hb*候选强行加入模型。",
    13,
)
replace_text(prs.slides[6].shapes[12],
    "13个位点的CatMAP稳态TOF已完成通量与稳态复核，并用于官方描述符图。\n\nTi–Zr没有通过验证，保留为未解决，不填TOF=0。\n\nTi–Re方向速率导出差异仍作为记录，不改变本轮TOF。",
    13,
)
replace_text(prs.slides[6].shapes[17],
    "第1、3、5、6步采用CatMAP官方端点势垒近似；第2、4步使用显式NEB过渡态。\n\n因此，ΔG3–TOF关系部分由端点势垒假设直接决定，不是独立的普适描述符验证。\n\n覆盖度或单一步骤能垒均不足以单独证明决速步骤。",
    13,
)
replace_text(
    prs.slides[6].shapes[18],
    "结论：当前数据支持统一模型内的机理比较；对真实催化活性及普适描述符的结论仍需补足关键步骤势垒并验证模型假设。",
    14,
)

# Slide 8: convert the obsolete “future recalculate everything” plan into a current roadmap.
replace_text(prs.slides[7].shapes[1], "从探索性CatMAP结果走向可检验的机理描述符", 23)
replace_text(
    prs.slides[7].shapes[2],
    "保持原始DFT、六步网络和333.15 K条件；后续只推进有证据支撑的机理验证。",
    14,
)
replace_text(prs.slides[7].shapes[6], "保持已验证模型边界")
replace_text(
    prs.slides[7].shapes[7],
    "六步反应网络与完整共吸附态定义保持不变；区分NEB显式TS与CatMAP端点势垒近似。",
    14,
)
replace_text(prs.slides[7].shapes[11], "检验电子结构—氧转移联系")
replace_text(
    prs.slides[7].shapes[12],
    "先整理已有Ti–M电子结构数据，检验其与NEB1氧转移势垒及稳态响应的关系；不预设相关方向或火山峰。",
    14,
)
replace_text(prs.slides[7].shapes[16], "补强关键步骤证据")
replace_text(
    prs.slides[7].shapes[17],
    "优先识别步骤1、3、5、6端点近似对结论的影响；只有获得或批准相应过渡态数据后，才更新正式动力学判断。",
    14,
)
replace_text(prs.slides[7].shapes[21], "维持结果分级与可追溯性")
replace_text(
    prs.slides[7].shapes[22],
    "13个位点保留探索性结果；Ti–Zr保持未解决。所有图继续显示已验证催化剂坐标和名称，不把失败位点补成零活性。",
    14,
)
replace_text(
    prs.slides[7].shapes[23],
    "研究目标：把模型内的速率响应与独立的电子结构、NEB势垒证据连接起来；不以拟合出火山关系作为成功标准。",
    14,
)

prs.save(str(DECK))
print(f"SAVED {DECK}")
