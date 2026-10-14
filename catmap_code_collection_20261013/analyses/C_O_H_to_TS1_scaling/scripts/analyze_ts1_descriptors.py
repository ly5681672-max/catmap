from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import shutil
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import load_workbook


SURFACES = [
    ("tife", "Ti-Fe"), ("timn", "Ti-Mn"), ("tihf", "Ti-Hf"),
    ("tire", "Ti-Re"), ("timo", "Ti-Mo"), ("tiv", "Ti-V"),
    ("tico", "Ti-Co"), ("titi", "Ti-Ti"), ("tiw", "Ti-W"),
    ("tita", "Ti-Ta"), ("ticr", "Ti-Cr"), ("ti", "Ti"),
    ("tini", "Ti-Ni"),
]
SURFACE_NAMES = {s for s, _ in SURFACES}
LABELS = {"C", "C*", "C_111", "C*_111", "C_atom_111", "C_atom", "C_s", "C_ads", "Cstar", "C_star", "C-atom", "C atom", "atomic C", "atomic carbon"}
SINGLE_MODELS = [("O",), ("Ha",), ("Hb",)]
PAIR_MODELS = [("O", "Ha"), ("O", "Hb"), ("Ha", "Hb")]


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def fit_ols(x: np.ndarray, y: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
    design = np.column_stack([x, np.ones(len(x))])
    coef, *_ = np.linalg.lstsq(design, y, rcond=None)
    return coef, design @ coef


def scores(y: np.ndarray, pred: np.ndarray) -> dict[str, float]:
    resid = pred - y
    denom = float(np.sum((y - np.mean(y)) ** 2))
    return {
        "R2_in_sample": float(1 - np.sum(resid ** 2) / denom) if denom > 0 else float("nan"),
        "RMSE_eV": float(np.sqrt(np.mean(resid ** 2))),
        "MAE_eV": float(np.mean(np.abs(resid))),
        "max_abs_error_eV": float(np.max(np.abs(resid))),
    }


def loocv_rmse(x: np.ndarray, y: np.ndarray) -> float:
    pred = np.empty_like(y, dtype=float)
    for i in range(len(y)):
        keep = np.arange(len(y)) != i
        coef, _ = fit_ols(x[keep], y[keep])
        pred[i] = float(np.r_[x[i], 1.0] @ coef)
    return float(np.sqrt(np.mean((pred - y) ** 2)))


def fit_model(df: pd.DataFrame, features: tuple[str, ...], actual_col: str) -> tuple[dict, pd.DataFrame]:
    x = df.loc[:, list(features)].to_numpy(dtype=float)
    y = df[actual_col].to_numpy(dtype=float)
    coef, pred = fit_ols(x, y)
    metrics = scores(y, pred)
    metrics.update({
        "features": "+".join(features), "N": len(df),
        "LOOCV_RMSE_eV": loocv_rmse(x, y),
        "intercept_eV": float(coef[-1]),
        "coefficients": json.dumps({f: float(v) for f, v in zip(features, coef[:-1])}, ensure_ascii=False),
        "fit_equation": " + ".join([f"({coef[i]:.12g})*{f}" for i, f in enumerate(features)] + [f"({coef[-1]:.12g})"]),
    })
    output = df[["surface_name", "catalyst", actual_col]].copy()
    for feature in features:
        output[feature + "_eV"] = df[feature].to_numpy()
    output["model"] = "+".join(features)
    output["predicted_TS1_eV"] = pred
    output["residual_pred_minus_actual_eV"] = pred - y
    return metrics, output


def scan_workbook(path: Path, cstar_rows: list[dict]) -> None:
    if not path.exists():
        cstar_rows.append({"source_path": str(path), "source_type": "workbook", "status": "NOT_FOUND"})
        return
    try:
        wb = load_workbook(path, read_only=True, data_only=False)
        sheets = list(wb.sheetnames)
        hits = []
        normalized = {label.lower() for label in LABELS}
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    value = cell.value
                    if isinstance(value, str) and value.strip().lower() in normalized:
                        hits.append((ws.title, cell.coordinate, value.strip()))
        wb.close()
        cstar_rows.append({
            "source_path": str(path), "source_type": "workbook_all_worksheets",
            "sha256": sha256(path), "sheet_count": len(sheets),
            "explicit_Cstar_label_count": len(hits),
            "matches": json.dumps(hits, ensure_ascii=False),
            "status": "EXPLICIT_CSTAR_LABELS_FOUND" if hits else "NO_EXPLICIT_CSTAR_LABELS",
        })
    except Exception as exc:
        cstar_rows.append({"source_path": str(path), "source_type": "workbook_all_worksheets", "status": f"READ_ERROR:{type(exc).__name__}:{exc}"})


def write_template(path: Path) -> None:
    fields = [
        "surface_name", "catalyst", "Cstar_structure_id", "Cstar_geometry_source",
        "Cstar_source_file", "Cstar_source_sheet", "Cstar_source_cell",
        "Cstar_total_G_eV", "Cstar_temperature_K", "Cstar_thermo_correction_eV",
        "clean_slab_structure_id", "clean_slab_source_file", "clean_slab_source_cell",
        "clean_slab_total_G_eV", "clean_slab_temperature_K", "clean_slab_thermo_correction_eV",
        "mu_C_eV", "mu_C_source_and_method", "functional_pseudopotential_cutoff_cell_kpoints",
        "spin_charge_settings", "formation_energy_formula", "Cstar_formation_energy_eV", "status",
    ]
    with path.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        for surface, catalyst in SURFACES:
            row = {k: "" for k in fields}
            row.update({
                "surface_name": surface, "catalyst": catalyst,
                "formation_energy_formula": "G(slab+C*) - G(clean slab) - mu_C",
                "status": "NEEDS_REAL_CSTAR_DFT_AND_REFERENCE_PROVENANCE",
            })
            writer.writerow(row)


def save_source_inventory(path: Path, rows: list[dict]) -> None:
    keys = sorted({key for row in rows for key in row})
    with path.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=keys)
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--project-root", type=Path, required=True)
    parser.add_argument("--raw-project-root", type=Path, required=True)
    parser.add_argument("--workbook", type=Path, action="append", required=True)
    args = parser.parse_args()

    out_root = Path(__file__).resolve().parents[1]
    paths = {name: out_root / name for name in ("data", "figures", "origin_exports", "reports")}
    for path in paths.values():
        path.mkdir(parents=True, exist_ok=True)
    data_dir, fig_dir, origin_dir = paths["data"], paths["figures"], paths["origin_exports"]
    source_data = args.project_root / "analyses" / "empirical_O_H_descriptor_analysis" / "data"
    energy_path = args.project_root / "catmap_official_run" / "energies.txt"
    formation_path = args.raw_project_root / "2-creating_microkinetic_model" / "DFT_formation_catmap.txt"
    converter_path = args.raw_project_root / "1-generating_input_file" / "convert_dft_to_catmap.py"
    if not all(p.is_file() for p in [source_data / "matched_data.csv", source_data / "workbook_source_audit.csv", source_data / "reference_slab_comparison.csv", energy_path, formation_path, converter_path]):
        raise FileNotFoundError("One or more required, read-only source artifacts are missing")

    # Preserve exact source snapshots in this analysis package.
    for source in [source_data / "matched_data.csv", source_data / "workbook_source_audit.csv", source_data / "reference_slab_comparison.csv", energy_path, formation_path, converter_path]:
        target = data_dir / ("source_" + source.name)
        if not target.exists():
            shutil.copy2(source, target)

    cstar_rows: list[dict] = []
    for workbook in args.workbook:
        scan_workbook(workbook, cstar_rows)
    source_file_candidates = []
    for base in [args.raw_project_root, args.project_root]:
        for p in base.rglob("*"):
            if not p.is_file():
                continue
            text = str(p).lower()
            if re.search(r"(cstar|c_star|c_atom|c_ads|atomic.?carbon|carbon.?ads)", text):
                source_file_candidates.append(str(p))
    cstar_rows.append({
        "source_path": str(args.raw_project_root), "source_type": "filename_scan",
        "matches": json.dumps(sorted(set(source_file_candidates)), ensure_ascii=False),
        "status": "CSTAR_NAMED_FILES_FOUND" if source_file_candidates else "NO_CSTAR_NAMED_COMPUTATIONAL_FILES",
    })

    formation = pd.read_csv(formation_path, sep="\t")
    cstar_species = formation[formation["species_name"].astype(str).str.strip().str.lower().isin({"c", "c*", "c_111", "c*_111", "c_atom", "c_atom_111", "c_s", "c_ads", "cstar", "c_star"})]
    cstar_rows.append({
        "source_path": str(formation_path), "source_type": "legacy_catmap_formation_table",
        "sha256": sha256(formation_path), "species_seen": json.dumps(sorted(formation["species_name"].astype(str).unique()), ensure_ascii=False),
        "explicit_Cstar_species_count": len(cstar_species),
        "status": "CSTAR_SPECIES_FOUND" if len(cstar_species) else "NO_CSTAR_SPECIES_IN_LEGACY_TABLE",
    })
    save_source_inventory(data_dir / "Cstar_source_inventory.csv", cstar_rows)
    write_template(data_dir / "Cstar_13site_DFT_data_template.csv")

    matched = pd.read_csv(source_data / "matched_data.csv")
    matched["surface_name"] = matched["surface_name"].astype(str).str.strip().str.lower()
    if matched["surface_name"].duplicated().any():
        raise ValueError("Duplicate surface_name in O/H/TOF matched data")
    if set(matched["surface_name"]) != SURFACE_NAMES:
        raise ValueError(f"Expected exactly 13 non-TiZr surfaces; got {sorted(matched.surface_name)}")
    if not np.allclose(matched["temperature_K"].to_numpy(float), 333.15, rtol=0, atol=1e-8):
        raise ValueError("Descriptor dataset is not entirely at 333.15 K")
    if not matched["status"].astype(str).str.upper().eq("CONVERGED").all():
        raise ValueError("Matched TOF records include non-converged entries")

    energy_rows = []
    with energy_path.open(encoding="utf-8-sig") as stream:
        header = next(stream).rstrip("\n").split("\t")
        for line in stream:
            row = line.rstrip("\n").split("\t")
            if len(row) < 4 or row[0] == "None":
                continue
            if row[2] == "TS1" and row[0].strip().lower() in SURFACE_NAMES:
                energy_rows.append({"surface_name": row[0].strip().lower(), "TS1_formation_energy_eV": float(row[3]), "reference": row[5] if len(row) > 5 else ""})
    ts1 = pd.DataFrame(energy_rows)
    if ts1["surface_name"].duplicated().any():
        raise ValueError("Duplicate TS1 surface_name in energies.txt")
    if set(ts1.surface_name) != SURFACE_NAMES:
        raise ValueError(f"TS1 surfaces mismatch: {sorted(ts1.surface_name)}")
    data = matched.merge(ts1, how="inner", on="surface_name", validate="one_to_one")
    order = {surface: i for i, (surface, _) in enumerate(SURFACES)}
    data["_order"] = data.surface_name.map(order)
    data = data.sort_values("_order").drop(columns="_order").reset_index(drop=True)
    data["TS1_source_file"] = str(energy_path)
    data["descriptor_source_file"] = str(source_data / "workbook_source_audit.csv")
    data.to_csv(data_dir / "matched_TS1_O_H_TOF_data.csv", index=False, float_format="%.12g", encoding="utf-8-sig")

    has_cstar = len(cstar_species) > 0 or any(row.get("status") == "EXPLICIT_CSTAR_LABELS_FOUND" for row in cstar_rows)
    available_models = SINGLE_MODELS + PAIR_MODELS
    metric_rows, prediction_frames = [], []
    for features in available_models:
        metrics, predictions = fit_model(data, features, "TS1_formation_energy_eV")
        metrics["data_scope"] = "13_site_common_O_H_subset"
        metrics["Cstar_source_status"] = "AVAILABLE" if has_cstar else "UNAVAILABLE"
        metric_rows.append(metrics)
        prediction_frames.append(predictions)
    if has_cstar:
        # Never guess the schema or infer energies from other adsorbates; add C* only after provenance audit.
        cstar_data_path = data_dir / "verified_Cstar_formation_energy.csv"
        if not cstar_data_path.exists():
            raise RuntimeError("A C* label was found but verified C* energies have not been manually source-validated; no C fit will be fabricated")
        c_data = pd.read_csv(cstar_data_path)
        if "surface_name" not in c_data or "Cstar" not in c_data:
            raise ValueError("Verified C* table must contain surface_name,Cstar")
        c_data["surface_name"] = c_data.surface_name.astype(str).str.strip().str.lower()
        if c_data.surface_name.duplicated().any():
            raise ValueError("Duplicate C* surface_name")
        data_c = data.merge(c_data[["surface_name", "Cstar"]], on="surface_name", how="inner", validate="one_to_one")
        if len(data_c) < 3:
            raise ValueError("Insufficient common sites for C* regressions")
        for features in [("Cstar",), ("Cstar", "O"), ("Cstar", "Ha"), ("Cstar", "Hb")]:
            metrics, predictions = fit_model(data_c, features, "TS1_formation_energy_eV")
            metrics["data_scope"] = "common_Cstar_O_H_subset"
            metrics["Cstar_source_status"] = "AVAILABLE_SOURCE_AUDITED"
            metric_rows.append(metrics)
            prediction_frames.append(predictions)
    metrics_df = pd.DataFrame(metric_rows).sort_values(["RMSE_eV", "features"]).reset_index(drop=True)
    metrics_df.to_csv(data_dir / "descriptor_fit_metrics.csv", index=False, float_format="%.12g", encoding="utf-8-sig")
    predictions = pd.concat(prediction_frames, ignore_index=True)
    predictions.to_csv(data_dir / "descriptor_fit_predictions.csv", index=False, float_format="%.12g", encoding="utf-8-sig")
    availability = pd.DataFrame([
        {"candidate": "Cstar", "status": "AVAILABLE" if has_cstar else "NO_VERIFIED_INDEPENDENT_CSTAR"},
        *[{"candidate": "+".join(pair), "status": "FITTED"} for pair in PAIR_MODELS],
    ])
    if not has_cstar:
        availability = pd.concat([availability, pd.DataFrame([{"candidate": "Cstar+O", "status": "NOT_FIT_CSTAR_MISSING"}, {"candidate": "Cstar+Ha", "status": "NOT_FIT_CSTAR_MISSING"}, {"candidate": "Cstar+Hb", "status": "NOT_FIT_CSTAR_MISSING"}])], ignore_index=True)
    availability.to_csv(data_dir / "descriptor_pair_availability.csv", index=False, encoding="utf-8-sig")

    # Comparison figure: only fits with real descriptor values are plotted.
    fig, axes = plt.subplots(1, 2, figsize=(13, 5.4), constrained_layout=True)
    pair_metrics = metrics_df[metrics_df.features.isin(["O+Ha", "O+Hb", "Ha+Hb"])].copy()
    colors = {"O+Ha": "#2878B5", "O+Hb": "#E07A2D", "Ha+Hb": "#4C956C"}
    for ax, col, label in [(axes[0], "R2_in_sample", "In-sample $R^2$"), (axes[1], "RMSE_eV", "In-sample RMSE (eV)")]:
        bars = ax.bar(pair_metrics.features, pair_metrics[col], color=[colors[k] for k in pair_metrics.features])
        ax.set_ylabel(label)
        ax.set_xlabel("Real historical descriptors")
        ax.grid(axis="y", alpha=0.25)
        for bar, value in zip(bars, pair_metrics[col]):
            ax.text(bar.get_x() + bar.get_width()/2, bar.get_height(), f"{value:.3f}", ha="center", va="bottom", fontsize=9)
    fig.suptitle("TS1 formation-energy reconstruction from available O/H descriptors\nEmpirical OLS; historical O/H reference retained", fontsize=12)
    fig.savefig(fig_dir / "descriptor_pair_fit_performance.png", dpi=220)
    fig.savefig(fig_dir / "descriptor_pair_fit_performance.pdf")
    plt.close(fig)

    for feature_set in PAIR_MODELS:
        model = "+".join(feature_set)
        subset = predictions[predictions.model == model].copy()
        fig, ax = plt.subplots(figsize=(6.4, 6.0), constrained_layout=True)
        vals = np.concatenate([subset.TS1_formation_energy_eV.to_numpy(), subset.predicted_TS1_eV.to_numpy()])
        lo, hi = float(np.min(vals)), float(np.max(vals))
        pad = max((hi-lo)*0.06, 0.05)
        ax.plot([lo-pad, hi+pad], [lo-pad, hi+pad], "k--", lw=1, label="y=x")
        ax.scatter(subset.TS1_formation_energy_eV, subset.predicted_TS1_eV, s=42, color=colors[model], edgecolor="white", zorder=3, label="13 matched catalysts")
        for row in subset.itertuples():
            ax.annotate(row.catalyst, (row.TS1_formation_energy_eV, row.predicted_TS1_eV), xytext=(4, 3), textcoords="offset points", fontsize=7)
        metric = pair_metrics[pair_metrics.features == model].iloc[0]
        ax.set_xlim(lo-pad, hi+pad); ax.set_ylim(lo-pad, hi+pad)
        ax.set_xlabel("Actual TS1 formation energy (eV)")
        ax.set_ylabel("Predicted TS1 formation energy (eV)")
        ax.set_title(f"{model} → TS1\nN={int(metric.N)}, R²={metric.R2_in_sample:.3f}, RMSE={metric.RMSE_eV:.3f} eV")
        ax.legend(frameon=False)
        ax.grid(alpha=0.2)
        stem = model.replace("+", "_") + "_TS1_actual_vs_predicted"
        fig.savefig(fig_dir / f"{stem}.png", dpi=220)
        fig.savefig(fig_dir / f"{stem}.pdf")
        plt.close(fig)

    # Origin-friendly coefficients and points are distinct from CatMAP grid results.
    metrics_df.to_csv(origin_dir / "fit_coefficients_and_metrics.csv", index=False, float_format="%.12g", encoding="utf-8-sig")
    data[["surface_name", "catalyst", "O", "Ha", "Hb", "TS1_formation_energy_eV", "TOF", "log10_TOF", "temperature_K", "status"]].to_csv(origin_dir / "catalyst_points.csv", index=False, float_format="%.12g", encoding="utf-8-sig")
    predictions.to_csv(origin_dir / "predictions_residuals.csv", index=False, float_format="%.12g", encoding="utf-8-sig")

    # Reproducibility manifest: source hashes and generated file hashes, excluding itself.
    hash_rows = []
    for p in [source_data / "matched_data.csv", source_data / "workbook_source_audit.csv", source_data / "reference_slab_comparison.csv", energy_path, formation_path, converter_path, *args.workbook]:
        hash_rows.append({"role": "read_only_input", "path": str(p), "sha256": sha256(p) if p.is_file() else "MISSING"})
    for p in sorted(out_root.rglob("*")):
        if p.is_file() and p.name != "SHA256SUMS.csv":
            hash_rows.append({"role": "analysis_output", "path": str(p), "sha256": sha256(p)})
    pd.DataFrame(hash_rows).to_csv(data_dir / "SHA256SUMS.csv", index=False, encoding="utf-8-sig")
    print(json.dumps({"out_root": str(out_root), "Cstar_source_status": "REAL_CSTAR_NOT_FOUND" if not has_cstar else "CSTAR_REQUIRES_MANUAL_SOURCE_VALIDATION", "matched_sites": len(data), "metrics": metrics_df.to_dict(orient="records")}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
