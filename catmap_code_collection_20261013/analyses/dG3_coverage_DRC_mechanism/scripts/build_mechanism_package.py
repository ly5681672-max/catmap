"""Assemble CatMAP mechanism analysis from cached results only.

Reads immutable strict-baseline tables and the official CatMAP TS1/TS2 cache.
It never calls ReactionModel.run() or a mapper/solver.
"""
from __future__ import annotations

import csv
import hashlib
import json
import math
import os
import shutil
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from scipy.stats import pearsonr, spearmanr
from openpyxl import load_workbook

ROOT = Path(r"C:\Users\Administrator\Documents\codex探索\catmap_neb_coadsorption_model")
ANALYSIS = ROOT / "analyses" / "dG3_coverage_DRC_mechanism"
PREV = ROOT / "analyses" / "microkinetic_mechanism_13sites"
DATA = PREV / "data"
FIG = ANALYSIS / "figures"
EXP = ANALYSIS / "origin_exports"
OUTDATA = ANALYSIS / "data"
REPORTS = ANALYSIS / "reports"
ARCHIVE = Path(r"H:\catmap\catmap\catmap_code_collection_20261013")
ARCHIVE_ENERGIES = ARCHIVE / "catmap_official_run" / "energies.txt"
GRID_DIR = ARCHIVE / "catmap_official_run" / "TS1_TS2_descriptor_grid"
GRID_CACHE = GRID_DIR / "runs" / "grid_50x50" / "catmap_ts1_ts2_grid_50x50.pkl"
WB_AUTH = Path(r"H:\Doctoral_research\己烯环氧化.xlsx")
WB_PROJECT = Path(r"H:\catmap\catmap\lyq_alkene_epoxidation\己烯环氧化.xlsx")

LABELS = {
    "ti": "Ti", "tico": "Ti–Co", "tini": "Ti–Ni", "tife": "Ti–Fe",
    "tire": "Ti–Re", "timo": "Ti–Mo", "timn": "Ti–Mn", "ticr": "Ti–Cr",
    "tiw": "Ti–W", "tiv": "Ti–V", "tita": "Ti–Ta", "tihf": "Ti–Hf", "titi": "Ti–Ti",
}
ADS = ["Rpair_s", "TS1_s", "EpoxideHHO_s", "HHO_s", "TS2_s", "OHHa_s", "H2O_s"]
K_B_EV = 8.617333262145e-5
T_K = 333.15

for d in (FIG, EXP, OUTDATA, REPORTS):
    d.mkdir(parents=True, exist_ok=True)

plt.rcParams.update({"font.family": "Microsoft YaHei", "axes.unicode_minus": False,
                     "figure.dpi": 150, "savefig.dpi": 300})


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def write_df(df: pd.DataFrame, path: Path) -> None:
    df.to_csv(path, index=False, encoding="utf-8-sig", float_format="%.16g")


def read_energy_table(path: Path) -> pd.DataFrame:
    df = pd.read_csv(path, sep="\t", dtype=str, keep_default_na=False)
    req = {"surface_name", "species_name", "formation_energy"}
    if not req.issubset(df.columns):
        raise ValueError(f"energies file missing fields {req - set(df.columns)}")
    df["formation_energy"] = pd.to_numeric(df["formation_energy"], errors="raise")
    return df


def find_energy(energy: pd.DataFrame, surface: str, species: str) -> float:
    rows = energy[(energy.surface_name == surface) & (energy.species_name == species)]
    # CatMAP TableParser stores gas references once with an empty surface_name.
    if len(rows) == 0:
        global_rows = energy[energy.surface_name.astype(str).str.lower().isin(["", "none", "nan"])]
        rows = global_rows[global_rows.species_name == species]
    if len(rows) != 1:
        raise ValueError(f"Expected one {surface}/{species} energy row, got {len(rows)}")
    return float(rows.iloc[0].formation_energy)


def no_data_rerun_guard() -> None:
    if not GRID_CACHE.is_file():
        raise FileNotFoundError(f"Existing official CatMAP cache missing: {GRID_CACHE}")
    if not ARCHIVE_ENERGIES.is_file():
        raise FileNotFoundError(ARCHIVE_ENERGIES)


def make_official_grid_energy_profiles(energy: pd.DataFrame) -> tuple[pd.DataFrame, dict]:
    """Use official 0.3.1 scaler and ReactionModel energy API at 13 anchors + 5 cached nodes.

    This is energy evaluation only. The saved CatMAP turnover/coverage maps are read;
    no mapper or steady-state solver is run.
    """
    import catmap
    from catmap import ReactionModel
    if catmap.__version__ != "0.3.1":
        raise RuntimeError(f"Expected CatMAP 0.3.1, found {catmap.__version__}")
    old_cwd = Path.cwd()
    os.chdir(GRID_DIR)
    try:
        model = ReactionModel(setup_file="ts1_ts2_descriptor_grid.mkm")
        model.data_file = str(Path("runs") / "grid_50x50" / GRID_CACHE.name)
        model.load_data_file(overwrite=True)
        if model.turnover_frequency_map is None or len(model.turnover_frequency_map) != 2500:
            raise ValueError("Existing official cache did not load 2500 TOF nodes")
        model.parse_elementary_rxns(model.rxn_expressions)
        mechanisms = model.elementary_rxns
        cached = pd.read_csv(DATA / "ts1_ts2_grid_peak_and_sides.csv")
        profile_records = []
        anchor_records = []
        # At each real catalyst anchor, compare official scaler+gas correction+reaction
        # energy API against the already accepted strict-baseline deltas/barriers.
        effective = pd.read_csv(DATA / "reaction_energy_13sites.csv")
        barriers = pd.read_csv(DATA / "effective_barriers_13sites.csv")
        for surface in LABELS:
            xy = [find_energy(energy, surface, "TS1"), find_energy(energy, surface, "TS2")]
            energies = model.scaler.get_free_energies(xy)
            pressures = list(model.gas_pressures)
            gas_pressures = dict(zip(model.gas_names, pressures))
            corrected = dict(energies)
            for gas, pressure in gas_pressures.items():
                if pressure > 0:
                    corrected[gas] += K_B_EV * T_K * math.log(float(pressure))
            for i, rxn in enumerate(mechanisms, start=1):
                dg, ea = model.get_rxn_energy(rxn, corrected)
                ref = effective[(effective.surface == surface) & (effective.step == i)]
                bref = barriers[(barriers.surface == surface) & (barriers.step == i)]
                if len(ref) != 1 or len(bref) != 1:
                    raise ValueError(f"Missing accepted reaction result for {surface}, step {i}")
                anchor_records.append({"surface": surface, "catalyst": LABELS[surface], "step": i,
                    "official_scaler_delta_G_eV": float(dg), "cached_delta_G_eV": float(ref.iloc[0].delta_G_eV),
                    "delta_G_difference_eV": float(dg) - float(ref.iloc[0].delta_G_eV),
                    "official_model_forward_barrier_eV": float(ea),
                    "cached_forward_barrier_eV": float(bref.iloc[0].forward_effective_barrier_eV),
                    "barrier_difference_eV": float(ea) - float(bref.iloc[0].forward_effective_barrier_eV)})
        for _, row in cached.iterrows():
            xy = [float(row.TS1_s_eV), float(row.TS2_s_eV)]
            energies = model.scaler.get_free_energies(xy)
            pressures = list(model.gas_pressures)
            gas_pressures = dict(zip(model.gas_names, pressures))
            corrected = dict(energies)
            for gas, pressure in gas_pressures.items():
                if pressure > 0:
                    corrected[gas] += K_B_EV * T_K * math.log(float(pressure))
            r = {"TS1_s_eV": xy[0], "TS2_s_eV": xy[1],
                 "grid_location_role": row.grid_location_role,
                 "CatMAP_log10_TOF_C6H12O": float(row.CatMAP_log10_TOF_C6H12O),
                 "TOF_C6H12O": 10.0 ** float(row.CatMAP_log10_TOF_C6H12O)}
            for sp in ADS:
                r[sp + "_formation_energy_eV"] = float(energies[sp])
            for i, rxn in enumerate(mechanisms, start=1):
                dg, ea = model.get_rxn_energy(rxn, corrected)
                r[f"dG{i}_effective_eV"] = float(dg)
                if i in (2, 4):
                    r[f"step{i}_forward_barrier_eV"] = float(ea)
            for sp in ("C6H12_g", "H2O2_g", "H2O_g", "C6H12O_g"):
                r[sp + "_pressure_used"] = float(gas_pressures[sp])
            profile_records.append(r)
        anchor_df = pd.DataFrame(anchor_records)
        if anchor_df.delta_G_difference_eV.abs().max() > 2e-6:
            raise RuntimeError("Official scaler anchor ΔG does not reproduce strict cache within 2e-6 eV")
        if anchor_df.barrier_difference_eV.abs().max() > 2e-6:
            raise RuntimeError("Official scaler anchor barrier does not reproduce strict cache within 2e-6 eV")
        return pd.DataFrame(profile_records), {
            "catmap_version": catmap.__version__, "cached_grid_nodes": len(model.turnover_frequency_map),
            "anchor_max_abs_dG_diff_eV": float(anchor_df.delta_G_difference_eV.abs().max()),
            "anchor_max_abs_forward_barrier_diff_eV": float(anchor_df.barrier_difference_eV.abs().max()),
            "anchor_checks": anchor_df,
        }
    finally:
        os.chdir(old_cwd)


def load_and_join() -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    energy = read_energy_table(ARCHIVE_ENERGIES)
    coverage_types = {k: str for k in ["EpoxideHHO_s", "H2O_s", "HHO_s", "OHHa_s", "Rpair_s", "s"]}
    cov = pd.read_csv(DATA / "coverage_13sites.csv", dtype=coverage_types)
    drc = pd.read_csv(DATA / "rate_control_matrix.csv")
    tof_all = pd.read_csv(DATA / "gas_turnover_frequencies_13sites.csv")
    tof = tof_all[tof_all.gas_species == "C6H12O_g"].copy()
    rxn = pd.read_csv(DATA / "reaction_energy_13sites.csv")
    barriers = pd.read_csv(DATA / "effective_barriers_13sites.csv")
    status_all = pd.read_csv(DATA / "coverage_site_balance.csv", dtype={"vacancy_coverage": str, "steady_state_residual": str, "site_balance_error": str})
    status = status_all[["surface", "status", "vacancy_coverage", "steady_state_residual", "site_balance_error"]]
    for frame, label in ((cov, "coverage"), (drc, "DRC"), (tof, "TOF"), (status, "status")):
        if frame.surface.duplicated().any() or len(frame) != 13:
            raise ValueError(f"{label} table must have exactly 13 unique surface_name rows")
    base = cov.merge(drc.drop(columns="catalyst"), on="surface", validate="one_to_one")
    base = base.merge(tof[["surface", "turnover_frequency", "log10_abs_turnover_frequency"]], on="surface", validate="one_to_one")
    base = base.merge(status, on="surface", validate="one_to_one")
    base["catalyst"] = base["surface"].map(LABELS)
    if base.catalyst.isna().any() or set(base.surface) != set(LABELS):
        raise ValueError("surface_name mapping does not match the 13 validated catalysts")
    base["log10_TOF"] = pd.to_numeric(base.log10_abs_turnover_frequency)
    base["TOF_s-1"] = pd.to_numeric(base.turnover_frequency)
    for surface in base.surface:
        for sp in ADS + ["C6H12O_g"]:
            base.loc[base.surface == surface, sp + "_formation_energy_eV"] = find_energy(energy, surface, sp.removesuffix("_g").removesuffix("_s"))
    # Standard-state reaction ΔG3 from the unchanged formation energy table.
    base["dG3_standard_eV"] = (base["HHO_s_formation_energy_eV"] + base["C6H12O_g_formation_energy_eV"] - base["EpoxideHHO_s_formation_energy_eV"])
    base["theta_EpoxideHHO"] = base["EpoxideHHO_s_x"]
    base["theta_HHO"] = base["HHO_s_x"]
    base["theta_OHHa"] = base["OHHa_s_x"]
    # Use the high-precision site-balance residual; the old wide table rounds
    # Ti-Fe's ~2.3e-17 vacancy to zero.
    base["theta_empty"] = base["vacancy_coverage"]
    base["DRC_EpoxideHHO"] = base["EpoxideHHO_s_y"]
    base["DRC_HHO"] = base["HHO_s_y"]
    base["DRC_OHHa"] = base["OHHa_s_y"]
    base["DRC_TS1"] = base["TS1_s"]
    base["DRC_TS2"] = base["TS2_s"]
    base["TiFe_audit_status"] = np.where(base.surface == "tife", "SOURCE_REFERENCE_DIFFERENCE_UNRESOLVED; dynamics internally reproduced", "NOT_APPLICABLE")
    # Long per-step data keep exact source values available for joins and sensitivity checks.
    rxn["catalyst"] = rxn.surface.map(LABELS)
    barriers["catalyst"] = barriers.surface.map(LABELS)
    return base.sort_values("surface").reset_index(drop=True), rxn, barriers, energy


def validate_tife_workbook(energy: pd.DataFrame) -> tuple[pd.DataFrame, dict]:
    if not WB_AUTH.exists() or not WB_PROJECT.exists():
        raise FileNotFoundError("One or both Ti-Fe source workbooks are missing")
    records = []
    for path in (WB_AUTH, WB_PROJECT):
        wbv = load_workbook(path, data_only=True, read_only=True)
        wbf = load_workbook(path, data_only=False, read_only=True)
        ws_v = wbv["吉布斯自由能汇总"]
        ws_f = wbf["吉布斯自由能汇总"]
        for cell, role in [("F4", "strict clean slab"), ("F12", "Rpair"), ("F16", "TS1"),
                           ("F20", "EpoxideHHO"), ("F24", "HHO"), ("F28", "TS2"),
                           ("F32", "OHHa"), ("F36", "H2O"), ("F41", "state10 C6H12O+H2O")]:
            records.append({"workbook": str(path), "workbook_sha256": sha256(path), "sheet": ws_v.title,
                "cell": cell, "role": role, "cached_value": ws_v[cell].value,
                "formula": ws_f[cell].value, "E0_cell": f"B{ws_v[cell].row}",
                "E0_eV": ws_v[f"B{ws_v[cell].row}"].value,
                "temperature_cell": f"C{ws_v[cell].row}", "temperature_K": ws_v[f"C{ws_v[cell].row}"].value,
                "correction_cell": f"E{ws_v[cell].row}", "correction_eV": ws_v[f"E{ws_v[cell].row}"].value,
                "calculation_path_cell": f"D{ws_v[cell].row}", "calculation_path": ws_v[f"D{ws_v[cell].row}"].value})
        # Historical clean slab is state table row51, not F41.
        for cell, role in [("G51", "historical single-species clean slab total G"),
                           ("C51", "historical clean slab E0"), ("F51", "historical clean slab correction"),
                           ("E51", "historical clean slab calculation path")]:
            ws_oldv = wbv["catmap数据集"]
            ws_oldf = wbf["catmap数据集"]
            records.append({"workbook": str(path), "workbook_sha256": sha256(path), "sheet": ws_oldv.title,
                "cell": cell, "role": role, "cached_value": ws_oldv[cell].value,
                "formula": ws_oldf[cell].value, "E0_cell": "C51", "E0_eV": ws_oldv["C51"].value,
                "temperature_cell": "D51", "temperature_K": ws_oldv["D51"].value,
                "correction_cell": "F51", "correction_eV": ws_oldv["F51"].value,
                "calculation_path_cell": "E51", "calculation_path": ws_oldv["E51"].value})
        wbv.close(); wbf.close()
    audit = pd.DataFrame(records)
    write_df(audit, OUTDATA / "tife_workbook_cell_audit.csv")
    # Source numeric comparison from authoritative workbook cached values.
    wba = load_workbook(WB_AUTH, data_only=True, read_only=True)
    ws = wba["吉布斯自由能汇总"]
    old = wba["catmap数据集"]
    strict_slab, historical_slab, state10 = float(ws["F4"].value), float(old["G51"].value), float(ws["F41"].value)
    expected = {"Rpair": "F12", "TS1": "F16", "EpoxideHHO": "F20", "HHO": "F24", "TS2": "F28", "OHHa": "F32", "H2O": "F36"}
    refs = {"Rpair": {"C": 6, "H": 14, "O": 2}, "TS1": {"C": 6, "H": 14, "O": 2},
            "EpoxideHHO": {"C": 6, "H": 14, "O": 2}, "HHO": {"H": 2, "O": 1},
            "TS2": {"H": 2, "O": 1}, "OHHa": {"H": 2, "O": 1}, "H2O": {"H": 2, "O": 1}}
    mu = {"C": -9.28, "H": -1.11, "O": -4.37}
    for sp, cell in expected.items():
        comp = refs[sp]
        formation = float(ws[cell].value) - strict_slab - sum(comp[e] * mu[e] for e in comp)
        model_sp = sp
        model_val = find_energy(energy, "tife", model_sp)
        if abs(formation - model_val) > 2e-6:
            raise RuntimeError(f"TiFe formation energy mismatch {sp}: Excel-> {formation}, energies.txt -> {model_val}")
    wba.close()
    wba = load_workbook(WB_AUTH, data_only=True, read_only=True)
    ws = wba["吉布斯自由能汇总"]
    old = wba["catmap数据集"]
    metadata = {"authoritative_workbook_sha256": sha256(WB_AUTH), "project_workbook_sha256": sha256(WB_PROJECT),
        "strict_slab_F4_eV": strict_slab, "historical_slab_G51_eV": historical_slab,
        "strict_minus_historical_slab_eV": strict_slab - historical_slab,
        "strict_minus_historical_E0_eV": float(ws["B4"].value) - float(old["C51"].value),
        "strict_minus_historical_correction_eV": float(ws["E4"].value) - float(old["F51"].value),
        "state10_F41_eV_not_a_slab": state10, "state10_minus_strict_slab_eV": state10 - strict_slab,
        "formation_energy_reconstruction": "7/7 tife network species reproduced from strict F4 and atomic references within 2e-6 eV",
        "temperature_metadata": "333.15 K explicit for F4, F12, F16, F20, F32; blank in C24/C28/C36; no inference/fill applied",
        "structure_equivalence": "UNRESOLVED: workbook records calculation-path strings, but the archived package contains no matching POSCAR/CONTCAR/OUTCAR evidence to establish same slab composition/structure/setup"}
    wba.close()
    return audit, metadata


def correlations(base: pd.DataFrame) -> dict:
    y = base.log10_TOF.to_numpy(float)
    result = {}
    for xname in ("dG3_standard_eV", "EpoxideHHO_s", "HHO_s", "OHHa_s", "s"):
        x = base[xname].to_numpy(float)
        result[xname] = {"pearson_r": float(pearsonr(x, y).statistic), "pearson_p": float(pearsonr(x, y).pvalue),
                         "spearman_rho": float(spearmanr(x, y).statistic), "spearman_p": float(spearmanr(x, y).pvalue)}
    return result


def draw_labeled_scatter(ax, x, y, labels, colors=None, annotate=True):
    colors = colors or ["#b94c3b" if "Fe" in label else "#245b82" for label in labels]
    ax.scatter(x, y, c=colors, s=44, edgecolors="white", linewidths=.5, zorder=3)
    if annotate:
        candidates = [(8,8),(-30,8),(8,-16),(-32,-16),(0,23),(0,-29),(38,0),(-43,0),
                      (40,20),(-45,20),(40,-21),(-45,-21),(12,38),(-12,-40),(55,8),(-58,-8)]
        fig = ax.figure
        fig.canvas.draw()
        renderer = fig.canvas.get_renderer()
        axes_box = ax.get_window_extent(renderer)
        used = []
        for xi, yi, lab in zip(x, y, labels):
            scored = []
            for off in candidates:
                trial = ax.annotate(lab, (xi, yi), xytext=off, textcoords="offset points", fontsize=7)
                fig.canvas.draw()
                box = trial.get_window_extent(renderer).expanded(1.05, 1.08)
                outside = box.x0 < axes_box.x0 or box.x1 > axes_box.x1 or box.y0 < axes_box.y0 or box.y1 > axes_box.y1
                overlap = sum(box.overlaps(b) for b in used)
                score = 1000 * outside + 100 * overlap + abs(off[0]) + abs(off[1])
                scored.append((score, off))
                trial.remove()
                if not outside and overlap == 0:
                    break
            off = min(scored, key=lambda item: item[0])[1]
            ann = ax.annotate(lab, (xi, yi), xytext=off, textcoords="offset points", fontsize=7, zorder=5,
                              arrowprops={"arrowstyle":"-", "lw":.35, "color":"#777"})
            fig.canvas.draw()
            used.append(ann.get_window_extent(renderer).expanded(1.05, 1.08))


def plot_outputs(base: pd.DataFrame, peak_data: pd.DataFrame, energy_profiles: pd.DataFrame) -> None:
    # Fig 1: joint ΔG3/coverage/activity relationships; preserves all 13 labels.
    fig, axes = plt.subplots(1, 3, figsize=(15.5, 4.7), constrained_layout=True)
    labels = base.catalyst.tolist()
    d3 = base.dG3_standard_eV.to_numpy(float)
    y = base.log10_TOF.to_numpy(float)
    draw_labeled_scatter(axes[0], d3, y, labels)
    theta_epox = pd.to_numeric(base.theta_EpoxideHHO).to_numpy(float)
    axes[0].set(xlabel="Delta G3 standard (eV)", ylabel="log10 TOF (s-1)", title="Delta G3 vs activity")
    axes[1].scatter(d3, theta_epox, c=["#b94c3b" if x == "tife" else "#245b82" for x in base.surface], edgecolors="white", s=44, zorder=3)
    saturated = base[pd.to_numeric(base.theta_EpoxideHHO) > .5].sort_values("dG3_standard_eV")
    label_y = np.linspace(.96, .12, len(saturated))
    for (_, r), ylab in zip(saturated.iterrows(), label_y):
        axes[1].annotate(r.catalyst, (r.dG3_standard_eV, float(r.theta_EpoxideHHO)),
                         xytext=(3.05, ylab), textcoords="data", fontsize=7,
                         arrowprops={"arrowstyle":"-", "lw":.35, "color":"#777"}, zorder=5)
    for _, r in base[pd.to_numeric(base.theta_EpoxideHHO) <= .5].iterrows():
        axes[1].annotate(r.catalyst, (r.dG3_standard_eV, float(r.theta_EpoxideHHO)), xytext=(5, 5),
                         textcoords="offset points", fontsize=7, zorder=5)
    axes[1].set_xlim(min(d3) - .08, 4.0)
    axes[1].set(xlabel="Delta G3 standard (eV)", ylabel="Coverage(EpoxideHHO*)", title="Delta G3 vs epoxide coverage")
    draw_labeled_scatter(axes[2], theta_epox, y, labels)
    axes[2].set(xlabel="Coverage(EpoxideHHO*)", ylabel="log10 TOF (s-1)", title="Epoxide coverage vs activity")
    fig.suptitle("Strict baseline: 13 validated sites; Ti-Fe highlighted", fontsize=12)
    fig.savefig(FIG / "dG3_coverage_TOF_joint.png", bbox_inches="tight"); fig.savefig(FIG / "dG3_coverage_TOF_joint.pdf", bbox_inches="tight"); plt.close(fig)

    # Fig 2: coverage + DRC paired view.
    drc_cols = ["DRC_EpoxideHHO", "DRC_HHO", "DRC_OHHa", "DRC_TS1", "DRC_TS2"]
    sub = base.set_index("catalyst").sort_values("catalyst")
    fig, axes = plt.subplots(1, 2, figsize=(14, 5.5), gridspec_kw={"width_ratios": [1.25, 1]}, constrained_layout=True)
    im = axes[0].imshow(sub[drc_cols].to_numpy(float), aspect="auto", cmap="coolwarm", vmin=-1.1, vmax=1.1)
    axes[0].set_xticks(range(len(drc_cols)), ["Epoxide", "HHO", "OHHa", "TS1", "TS2"], rotation=30, ha="right")
    axes[0].set_yticks(range(len(sub)), sub.index); axes[0].set_title("Official product DRC (local sensitivity)")
    fig.colorbar(im, ax=axes[0], label="DRC")
    x = np.arange(len(sub)); width = .25
    axes[1].bar(x-width, pd.to_numeric(sub.theta_EpoxideHHO), width, label="theta(EpoxideHHO)")
    axes[1].bar(x, pd.to_numeric(sub.theta_HHO), width, label="theta(HHO)")
    axes[1].bar(x+width, pd.to_numeric(sub.theta_OHHa), width, label="theta(OHHa)")
    axes[1].set_xticks(x, sub.index, rotation=45, ha="right"); axes[1].set_ylim(0, 1.05)
    axes[1].set_title("Steady-state coverages"); axes[1].legend(fontsize=8)
    fig.savefig(FIG / "coverage_DRC_13sites.png", bbox_inches="tight"); fig.savefig(FIG / "coverage_DRC_13sites.pdf", bbox_inches="tight"); plt.close(fig)

    # Fig 3: TOF + coverage changes on each cached descriptor-axis cross-section.
    fig, axes = plt.subplots(2, 2, figsize=(12, 8), constrained_layout=True)
    roles = ["TS1_lower_side", "grid_peak", "TS1_higher_side", "TS2_lower_side", "grid_peak", "TS2_higher_side"]
    for ax, direction, selection in [(axes[0,0], "TS1_s", roles[:3]), (axes[0,1], "TS2_s", roles[3:])]:
        d = peak_data[peak_data.grid_location_role.isin(selection)].copy()
        d = d.sort_values(direction + "_eV")
        xx = d[direction + "_eV"].to_numpy(float)
        ax.plot(xx, d.CatMAP_log10_TOF_C6H12O, "o-", label="log10 TOF")
        ax.set_xlabel(direction + " (eV)"); ax.set_ylabel("log10 TOF (s-1)"); ax.set_title(direction + " axis (other descriptor fixed at peak)")
        for _, r in d.iterrows(): ax.annotate(r.grid_location_role.replace("_", " "), (r[direction+"_eV"], r.CatMAP_log10_TOF_C6H12O), xytext=(3,5), textcoords="offset points", fontsize=7)
    for ax, direction, selection in [(axes[1,0], "TS1_s", roles[:3]), (axes[1,1], "TS2_s", roles[3:])]:
        d = peak_data[peak_data.grid_location_role.isin(selection)].copy().sort_values(direction + "_eV")
        ax.plot(d[direction+"_eV"], d.EpoxideHHO_s, "o-", label="EpoxideHHO")
        ax.plot(d[direction+"_eV"], d.H2O_s, "s-", label="H2O")
        ax.plot(d[direction+"_eV"], d.Rpair_s, "^-", label="Rpair")
        ax.set_xlabel(direction + " (eV)"); ax.set_ylabel("Coverage"); ax.set_ylim(-.03, 1.03); ax.set_title(direction + " axis coverage redistribution"); ax.legend(fontsize=8)
    fig.savefig(FIG / "TS1_TS2_peak_coverage_TOF.png", bbox_inches="tight"); fig.savefig(FIG / "TS1_TS2_peak_coverage_TOF.pdf", bbox_inches="tight"); plt.close(fig)

    # Fig 4: local energy control variables along the same existing cached nodes.
    fig, axes = plt.subplots(1, 2, figsize=(12, 4.8), constrained_layout=True)
    for ax, direction in zip(axes, ["TS1_s", "TS2_s"]):
        d = energy_profiles[energy_profiles.grid_location_role.isin(roles[:3] if direction == "TS1_s" else roles[3:])].copy().sort_values(direction + "_eV")
        ax.plot(d[direction+"_eV"], d.dG2_effective_eV, "o-", label="dG2 eff")
        ax.plot(d[direction+"_eV"], d.dG3_effective_eV, "s-", label="dG3 eff")
        ax.plot(d[direction+"_eV"], d.dG4_effective_eV, "^-", label="dG4 eff")
        ax.set_xlabel(direction + " (eV)"); ax.set_ylabel("Effective reaction ΔG (eV)"); ax.set_title(direction + "-axis effective reaction energies"); ax.legend(fontsize=8)
    fig.savefig(FIG / "TS1_TS2_peak_reaction_energy.png", bbox_inches="tight"); fig.savefig(FIG / "TS1_TS2_peak_reaction_energy.pdf", bbox_inches="tight"); plt.close(fig)

    # Fig 5: TiFe focused contrast with representative sites.
    chosen = base[base.surface.isin(["tife", "tico", "tini", "ti"])].copy()
    fig, axes = plt.subplots(1, 2, figsize=(11, 4.5), constrained_layout=True)
    x = np.arange(len(chosen)); w = .2
    for j, sp in enumerate(["theta_EpoxideHHO", "theta_HHO", "theta_OHHa", "theta_empty"]):
        axes[0].bar(x+(j-1.5)*w, pd.to_numeric(chosen[sp]), w, label=sp)
    axes[0].set_xticks(x, chosen.catalyst); axes[0].set_ylim(0, 1.05); axes[0].set_title("Steady-state coverage"); axes[0].legend(fontsize=8)
    for j, sp in enumerate(["DRC_EpoxideHHO", "DRC_HHO", "DRC_OHHa", "DRC_TS1", "DRC_TS2"]):
        axes[1].bar(x+(j-2)*w, chosen[sp], w, label=sp)
    axes[1].axhline(0, color="black", lw=.5); axes[1].set_xticks(x, chosen.catalyst); axes[1].set_title("Official DRC"); axes[1].legend(fontsize=8)
    fig.suptitle("Ti–Fe versus Ti, Ti–Co, Ti–Ni"); fig.savefig(FIG / "TiFe_mechanism_comparison.png", bbox_inches="tight"); fig.savefig(FIG / "TiFe_mechanism_comparison.pdf", bbox_inches="tight"); plt.close(fig)


def write_reports(base: pd.DataFrame, combined: pd.DataFrame, corr_df: pd.DataFrame,
                  peak_data: pd.DataFrame, energy_profiles: pd.DataFrame,
                  tife_meta: dict, scaler_audit: dict) -> None:
    def stat(subset: str, name: str, col: str) -> float:
        row = corr_df[(corr_df.subset == subset) & (corr_df.descriptor == name)]
        return float(row[col].iloc[0])
    c_dg_tof = stat("13_sites", "dG3_standard_eV", "spearman_rho")
    c_dg_epox = spearmanr(pd.to_numeric(base.dG3_standard_eV), pd.to_numeric(base.theta_EpoxideHHO)).statistic
    c_dg_hho = spearmanr(pd.to_numeric(base.dG3_standard_eV), pd.to_numeric(base.theta_HHO)).statistic
    c_dg_empty = spearmanr(pd.to_numeric(base.dG3_standard_eV), pd.to_numeric(base.theta_empty)).statistic
    c_epox_tof = spearmanr(pd.to_numeric(base.theta_EpoxideHHO), pd.to_numeric(base.log10_TOF)).statistic
    tife = combined[combined.surface == "tife"].iloc[0]
    drc_11 = combined[(combined.surface != "ti") & (combined.surface != "tife")]
    near_pair = ((pd.to_numeric(drc_11.DRC_HHO) > .9) & (pd.to_numeric(drc_11.DRC_EpoxideHHO) < -.9)).sum()
    peak = energy_profiles[energy_profiles.grid_location_role == "grid_peak"].iloc[0]
    peak_text = []
    for direction, subset in [("TS1", energy_profiles[energy_profiles.grid_location_role.isin(["TS1_lower_side", "grid_peak", "TS1_higher_side"])]),
                              ("TS2", energy_profiles[energy_profiles.grid_location_role.isin(["TS2_lower_side", "grid_peak", "TS2_higher_side"])])]:
        subset = subset.sort_values("TS1_s_eV" if direction == "TS1" else "TS2_s_eV")
        peak_text.append(f"- {direction}轴：log10TOF范围 {subset.CatMAP_log10_TOF_C6H12O.min():.6f} 至 {subset.CatMAP_log10_TOF_C6H12O.max():.6f}；ΔG2/ΔG3/ΔG4与正向显式TS有效势垒见`data/TS1_TS2_peak_mechanism_data.csv`。")
    reuse = """# 已有成果复用与复核审计

本轮只读复用CatMAP 0.3.1严格基准结果；未运行`ReactionModel.run()`、稳态求解、官方描述符网格或储库条件搜索。

| 任务项 | 状态 | 证据/处理 |
|---|---|---|
| 13位点 ΔG3–TOF | `VERIFIED_REUSE` | 由严格基准形成能独立重算标准ΔG3，并与已有step-3有效ΔG按surface_name逐一合并；旧的activity_analysis_round1关系已存在，本轮只作一致性复核。 |
| EpoxideHHO*/HHO*覆盖度 | `VERIFIED_REUSE` | 复用高精度coverage_13sites.csv；空位另取coverage_site_balance.csv，避免Ti–Fe被旧宽表舍入成0。 |
| 13位点官方DRC | `VERIFIED_REUSE` | 复用rate_control_matrix.csv；没有将DRC解释成唯一决速步。 |
| TS1–TS2峰值及四个轴向邻点 | `VERIFIED_REUSE` | 复用已存5点表及50×50 CatMAP缓存的覆盖度/TOF；未重新求解网格。 |
| 网格点ΔG2/ΔG3/ΔG4和正向势垒 | `NOT_DONE -> 本轮补齐` | 通过CatMAP 0.3.1 `GeneralizedLinearScaler.get_free_energies`与`ReactionModel.get_rxn_energy`对既有5个缓存坐标作能量评估，并用13个真实锚点与基线输出交叉核验。 |
| Ti–Fe clean slab历史/严格参考比较 | `NEEDS_CORRECTION` | 旧reference_slab_comparison.csv将F41列作alternate slab；实际F41是状态10（C6H12O+H2O），历史slab应为catmap数据集!G51。已在新审计中更正，不改旧文件。 |
| Ti–Fe历史与严格slab结构等价性 | `UNCERTAIN` | Excel路径可读，但本归档没有相应POSCAR/CONTCAR/OUTCAR来证明结构/组成/计算设置一致；不作数值校正。 |

数据来源和各文件SHA256见`data/source_manifest_sha256.csv`。输出不替换旧图或历史审计文件。
"""
    (REPORTS / "EXISTING_WORK_REUSE_AUDIT.md").write_text(reuse, encoding="utf-8")

    tife_report = f"""# Ti–Fe能量来源与机制专项审计

## 结论先行

Ti–Fe的严格CatMAP能量输入、覆盖度、TOF和DRC在已有输出内部可追溯且相互一致；目前没有证据证明Ti–Fe严格结果存在映射/数值错误。它确实表现为与多数双金属位点不同的稳态响应：OHHa*覆盖度高、OHHa*具有显著负DRC。但历史单物种slab与严格模型slab的约1.064927 eV差异尚未能由结构文件解释，所以“物理结构导致的机制差异”仍需保留为有证据限制的判断。

## 原始Ti–Fe状态和输入换算

严格工作簿`吉布斯自由能汇总`中clean slab为F4；七个网络态原始G分别来自F12/F16/F20/F24/F28/F32/F36。每项为工作簿缓存公式值（公式为该行E0与Gibbs修正相加）；完整公式、温度单元格、计算路径和两工作簿SHA256见`data/tife_workbook_cell_audit.csv`。

| 状态 | 原始G单元格 | G (eV) | 温度记录 |
|---|---:|---:|---|
| Rpair* | F12 | -984.519289 | 333.15 K |
| TS1* | F16 | -984.359873 | 333.15 K |
| EpoxideHHO* | F20 | -986.698272 | 333.15 K |
| HHO* | F24 | -884.582039 | 空白 |
| TS2* | F28 | -884.4090575 | 空白 |
| OHHa* | F32 | -885.199105 | 333.15 K |
| H2O* | F36 | -885.0369905 | 空白 |

按当前形成能参考（clean slab F4；μC=-9.28、μH=-1.11、μO=-4.37 eV/atom）逐态重构，7/7与当前CatMAP energies.txt相差不超过2e-6 eV。反应中间体组成仍采用模型定义：Rpair/TS1/EpoxideHHO为C6H14O2，HHO/TS2/OHHa/H2O为H2O总组成；HHO不重新解释成水分子结构。

## 约1.065 eV差异的真实来源

- 严格网络clean slab：`吉布斯自由能汇总!F4 = -869.897259 eV`，E0=-870.071242 eV，修正=+0.173983 eV。
- 历史单物种clean slab：`catmap数据集!G51 = -870.962186 eV`，E0=-871.128961 eV，修正=+0.166775 eV。
- 严格减历史 = +1.064927 eV；其中E0差 +1.057719 eV、热修正差 +0.007208 eV。
- 旧表中的严格“alternate slab”`F41=-870.960555 eV`实际是状态10，组标题为C6H12O+H2O，不是clean slab；不能用它替代G51。
- Excel计算路径分别指向`tife/tiofe-t/freq`、历史`.../tife/catmap/surface`等目录；归档中没有能证明相同结构/组成/设置的原始结构输出。差异根因因此仍是**来源与结构参照未闭合**，不是已证明的DFT错误。未将1.064927 eV加减回任一能量。

NEB2相关HHO/TS2/H2O的Excel温度单元格C24/C28/C36为空；没有在本轮推断温度。严格模型采用冻结Gibbs数据/333.15 K求解的既有约定，但源表本身的这些热修正温度证据不完整，仍属数据来源限制。

## 已有动力学复核

- θ(OHHa*)={float(tife.theta_OHHa):.9f}；θ(EpoxideHHO*)={float(tife.theta_EpoxideHHO):.9f}；空位高精度值={tife.theta_empty}；
- TOF={float(tife['TOF_s-1']):.9g} s-1，log10TOF={float(tife.log10_TOF):.6f}；严格稳态残差={tife.steady_state_residual}，位点守恒误差={tife.site_balance_error}；
- DRC(OHHa*)={float(tife.DRC_OHHa):.6f}，DRC(EpoxideHHO*)={float(tife.DRC_EpoxideHHO):.6f}，DRC(HHO*)={float(tife.DRC_HHO):.6f}；
- 标准ΔG3={float(tife.dG3_standard_eV):.6f} eV，有效ΔG3={float(tife.dG3_effective_eV):.6f} eV。

这些指标均从现有13位点严格输出按`surface_name=tife`重取；不重跑，不覆盖。Ti–Fe应保留在全体13点分析中，并可另做12位点排除敏感性对照，但不能因此把它当成错误点。
"""
    (REPORTS / "TIFE_DATA_AND_MECHANISM_AUDIT.md").write_text(tife_report, encoding="utf-8")

    tof_by_role = {r.grid_location_role: r for _, r in energy_profiles.iterrows()}
    center = peak_data[peak_data.grid_location_role == "grid_peak"].iloc[0]
    peak_section = "\n".join(peak_text)
    report = f"""# ΔG3–覆盖度–DRC–TOF及TS1–TS2峰值机制关联

## 数据与口径

使用CatMAP 0.3.1严格基准、333.15 K、13个已收敛位点；Ti–Zr未纳入。13位点数据按surface_name一对一合并。标准态ΔG3按形成能直接计算：

`ΔG3_standard = Gform(HHO*) + Gform(C6H12O_g) - Gform(EpoxideHHO*)`。

有效ΔG3取既有CatMAP固定储库严格基准的step-3结果；它包含气相化学势修正，与标准ΔG3的定义不同。两者在当前固定产物储库下相差约{float((combined.dG3_effective_eV-combined.dG3_standard_eV).mean()):.9f} eV且为常数，因此相关系数数值相同，不是两种独立证据。

## 13位点关联结果

- 标准ΔG3与log10TOF：Spearman ρ={c_dg_tof:.3f}，Pearson r={stat('13_sites','dG3_standard_eV','pearson_r'):.3f}（N=13）。这复现了此前严格基准的单调相关，但本身不证明ΔG3是唯一控速因素。
- 标准ΔG3与EpoxideHHO覆盖度：Spearman ρ={c_dg_epox:.3f}；与HHO覆盖度ρ={c_dg_hho:.3f}；与高精度空位覆盖度ρ={c_dg_empty:.3f}。
- EpoxideHHO覆盖度与log10TOF：Spearman ρ={c_epox_tof:.3f}。多个位点覆盖度极接近1，需结合完整数值和DRC，不把近饱和等同于唯一限速。
- 13位点中的多数双金属位点（不含Ti–Fe的11个）有{near_pair}/11同时满足DRC(HHO)>0.9和DRC(EpoxideHHO)<-0.9；这支持局部活性对两种能量的相反敏感性。DRC仍是局部导数，不是唯一决速步骤标签。
- Ti–Fe不同：θ(OHHa*)≈{float(tife.theta_OHHa):.3f}且DRC(OHHa*)≈{float(tife.DRC_OHHa):.3f}，EpoxideHHO覆盖仅≈{float(tife.theta_EpoxideHHO):.3f}。这与OHHa积累/后段敏感的特殊稳态响应一致；同时其clean-slab来源差异仍未物理闭合，故不把它提升为已验证的本征机制因果结论。
- Ti单核位点空位约{float(combined[combined.surface=='ti'].theta_empty.iloc[0]):.6g}，DRC(TS1)≈{float(combined[combined.surface=='ti'].DRC_TS1.iloc[0]):.3f}，与多数双金属位点不同。

因此较稳妥的解释是：在多数异核位点上，较不利的环氧产物脱附端点热力学与强EpoxideHHO稳定/覆盖、HHO/EpoxideHHO相反的DRC共同形成活性趋势；Ti–Fe及Ti表明还有OHHa或TS1敏感的不同响应。第3步采用CatMAP端点势垒近似，ΔG3–TOF关系部分由该模型定义耦合，不能作为独立验证的普适描述符规律。

## TS1–TS2既有CatMAP网格峰及两侧

既有官方50×50 CatMAP缓存给出离散峰TS1={float(center.TS1_s_eV):.6f} eV、TS2={float(center.TS2_s_eV):.6f} eV、log10TOF={float(center.CatMAP_log10_TOF_C6H12O):.6f}。本轮仅复用峰及其四个已缓存轴向邻点，并调用官方Scaler对这些坐标作能量读取；官方能量API在13个真实位点上与strict baseline最大差：ΔG {scaler_audit['anchor_max_abs_dG_diff_eV']:.3g} eV、正向势垒 {scaler_audit['anchor_max_abs_forward_barrier_diff_eV']:.3g} eV。

{peak_section}

峰处θ(EpoxideHHO)={float(center.EpoxideHHO_s):.3f}、θ(H2O)={float(center.H2O_s):.3f}、θ(Rpair)={float(center.Rpair_s):.3f}；邻点显示TOF下降伴随这三种状态重新分配，而非单一覆盖物种替换可以充分概括。TS1两侧ΔG2/forward barrier和TS2两侧ΔG3/ΔG4变化均列在`data/TS1_TS2_peak_mechanism_data.csv`。

这些虚拟坐标的物种能量来自CatMAP标度器，不是新增DFT；也没有缓存该峰附近的DRC，因此峰两侧的机制解释是能量—覆盖度—TOF的探索性关联，不能声称已由局部DRC直接验证。

## 全体/12位点敏感性与限制

`data/TIFE_13site_vs_12site_comparison.csv`保留全体13位点，并标记包含Ti–Fe与否；12位点结果只作敏感性对照。排除Ti–Fe后标准ΔG3–log10TOF Spearman ρ={stat('12_sites_excluding_TiFe_sensitivity','dG3_standard_eV','spearman_rho'):.3f}，与13位点趋势一致；这不构成排除Ti–Fe的理由。

图件包括13位点联合关系、覆盖度/DRC、峰值轴向覆盖度与TOF、峰值有效反应能，以及Ti–Fe代表位点对照；对应CSV均在`origin_exports/`。未重新绘制此前已有的单变量活动图，也未计算虚拟网格DRC。

## 科研表达分级

**可直接用于PPT（注明模型条件）：** 13位点严格基准的标准/有效ΔG3区分、覆盖度与TOF排序、官方DRC局部响应，以及CatMAP缓存峰/两侧的覆盖重分配；Ti–Fe具有不同的数值稳态响应。

**需补原始结构/计算证据后再作更强物理结论：** Ti–Fe历史与严格slab为何相差约1.065 eV；NEB2若干Gibbs修正的温度记录来源；峰区的DRC或更密的局部数据。上述缺口不影响本轮复用既有基线结果，但限制因果归因。
"""
    (REPORTS / "MECHANISM_CORRELATION_REPORT.md").write_text(report, encoding="utf-8")


def main() -> None:
    no_data_rerun_guard()
    base, rxn, barriers, energy = load_and_join()
    # Form ΔG3 from explicit CatMAP energies; compare to the already archived strict effective values.
    rxn3 = rxn[rxn.step == 3][["surface", "delta_G_eV"]].rename(columns={"delta_G_eV": "dG3_effective_eV"})
    base = base.merge(rxn3, on="surface", validate="one_to_one")
    # Energy audit and old-vs-strict source discrepancy.
    wb_audit, tife_meta = validate_tife_workbook(energy)
    if "theta_EpoxideHHO" not in base.columns:
        raise RuntimeError("Coverage/DRC merge lost expected species columns")
    base["TOF_source_status"] = "STRICT_BASELINE_CONVERGED_REUSED"
    combined_cols = ["surface", "catalyst", "status", "dG3_standard_eV", "dG3_effective_eV",
        "EpoxideHHO_s_formation_energy_eV", "HHO_s_formation_energy_eV", "theta_EpoxideHHO", "theta_HHO", "theta_OHHa", "theta_empty",
        "DRC_EpoxideHHO", "DRC_HHO", "DRC_OHHa", "DRC_TS1", "DRC_TS2", "TOF_s-1", "log10_TOF", "steady_state_residual", "site_balance_error", "TiFe_audit_status", "TOF_source_status"]
    combined = base[combined_cols].copy()
    write_df(combined, OUTDATA / "combined_dG3_coverage_DRC_TOF.csv")
    write_df(combined, EXP / "combined_dG3_coverage_DRC_TOF_origin.csv")

    # 12-site sensitivity is separate and never replaces the full 13-site dataset.
    combined["included_in_13site_full_set"] = True
    combined["included_in_12site_sensitivity"] = combined.surface != "tife"
    write_df(combined, OUTDATA / "TIFE_13site_vs_12site_comparison.csv")
    write_df(combined[combined.surface != "tife"], OUTDATA / "TIFE_12site_sensitivity_only.csv")
    # Produce an explicit compact correlation table with pairwise N.
    cor_rows = []
    for group, subset in [("13_sites", base), ("12_sites_excluding_TiFe_sensitivity", base[base.surface != "tife"])]:
        for xname in ["dG3_standard_eV", "dG3_effective_eV", "EpoxideHHO_s_x", "HHO_s_x", "OHHa_s_x", "s"]:
            if xname not in subset.columns: continue
            x = subset[xname].to_numpy(float); y = subset.log10_TOF.to_numpy(float)
            pr, pp = pearsonr(x, y); sr, sp = spearmanr(x, y)
            cor_rows.append({"subset": group, "descriptor": xname, "N": len(subset), "pearson_r": pr, "pearson_p": pp, "spearman_rho": sr, "spearman_p": sp})
    corr_df = pd.DataFrame(cor_rows)
    write_df(corr_df, OUTDATA / "dG3_coverage_TOF_correlations.csv")

    # Existing official grid: reuse the cached 5 representative peak/side points only.
    peak_data = pd.read_csv(DATA / "ts1_ts2_grid_peak_and_sides.csv")
    energy_profiles, scaler_audit = make_official_grid_energy_profiles(energy)
    write_df(energy_profiles, OUTDATA / "TS1_TS2_peak_mechanism_data.csv")
    write_df(energy_profiles, EXP / "TS1_TS2_peak_mechanism_data_origin.csv")
    write_df(scaler_audit["anchor_checks"], OUTDATA / "official_scaler_anchor_reproduction.csv")
    write_df(scaler_audit["anchor_checks"], EXP / "official_scaler_anchor_reproduction_origin.csv")
    # Peak/side data with coverage joined by coordinate-role identity (not row order).
    peak_data = peak_data.merge(energy_profiles, on=["TS1_s_eV", "TS2_s_eV", "grid_location_role", "CatMAP_log10_TOF_C6H12O"], validate="one_to_one", suffixes=("_coverage_cache", "_energy_eval"))
    write_df(peak_data, EXP / "TS1_TS2_peak_coverage_TOF_origin.csv")
    plot_outputs(base, peak_data, energy_profiles)

    # Preserve the original cached local coverages; no refitting or interpolation.
    pd.read_csv(DATA / "ts1_ts2_grid_coverage_2500nodes.csv").to_csv(EXP / "TS1_TS2_grid_coverage_2500nodes_origin.csv", index=False, encoding="utf-8-sig")
    # Short source manifest hashes only (no modifications to source files).
    source_paths = [ARCHIVE_ENERGIES, GRID_DIR / "ts1_ts2_descriptor_grid.mkm", GRID_CACHE,
                    WB_AUTH, WB_PROJECT, DATA / "reaction_energy_13sites.csv", DATA / "effective_barriers_13sites.csv",
                    DATA / "coverage_13sites.csv", DATA / "rate_control_matrix.csv", DATA / "gas_turnover_frequencies_13sites.csv",
                    DATA / "ts1_ts2_grid_peak_and_sides.csv", DATA / "coverage_site_balance.csv"]
    manifest = [{"source_path": str(p), "sha256": sha256(p), "bytes": p.stat().st_size} for p in source_paths]
    write_df(pd.DataFrame(manifest), OUTDATA / "source_manifest_sha256.csv")
    summary = {"sites": len(combined), "unique_surface_names": combined.surface.nunique(), "CatMAP_recalculated": False,
        "descriptor_grid_recalculated": False, "condition_search_recalculated": False,
        "tife": tife_meta,
        "official_scaler": {k: v for k, v in scaler_audit.items() if k != "anchor_checks"},
        "dG3_13site_spearman": float(corr_df[(corr_df.subset == "13_sites") & (corr_df.descriptor == "dG3_standard_eV")].spearman_rho.iloc[0]),
        "dG3_12site_sensitivity_spearman": float(corr_df[(corr_df.subset == "12_sites_excluding_TiFe_sensitivity") & (corr_df.descriptor == "dG3_standard_eV")].spearman_rho.iloc[0])}
    (OUTDATA / "analysis_validation_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    write_reports(base, combined, corr_df, peak_data, energy_profiles, tife_meta, scaler_audit)
    # Audit log is deliberately deterministic and records gating checks.
    with (REPORTS / "run_validation.log").open("w", encoding="utf-8") as f:
        f.write(f"CatMAP recalculated: {summary['CatMAP_recalculated']}\n")
        f.write(f"Descriptor grid recalculated: {summary['descriptor_grid_recalculated']}\n")
        f.write(f"Condition search recalculated: {summary['condition_search_recalculated']}\n")
        f.write(f"Joined sites: {len(combined)} / unique surfaces {combined.surface.nunique()}\n")
        f.write(f"Maximum official anchor |ΔG difference|: {scaler_audit['anchor_max_abs_dG_diff_eV']:.12g} eV\n")
        f.write(f"Maximum official anchor |forward barrier difference|: {scaler_audit['anchor_max_abs_forward_barrier_diff_eV']:.12g} eV\n")
        f.write(f"TiFe F4-G51: {tife_meta['strict_minus_historical_slab_eV']:.9f} eV; F41 is state10, not slab\n")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
