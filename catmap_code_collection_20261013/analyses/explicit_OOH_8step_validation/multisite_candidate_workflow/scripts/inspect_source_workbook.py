"""Read-only inspection of workbook cells needed for a non-TiCo candidate."""
from __future__ import annotations

import os
from pathlib import Path
import openpyxl

book_path = Path(os.environ["OOH_WORKBOOK_PATH"])
for data_only in (False, True):
    wb = openpyxl.load_workbook(book_path, data_only=data_only, read_only=True)
    ws = wb.worksheets[0]
    print(f"data_only={data_only} sheet={ws.title!r} shape={ws.max_row}x{ws.max_column}")
    for row in list(range(1, 22)) + list(range(48, 62)):
        cells = []
        for col in list(range(1, 8)) + list(range(60, 68)):
            value = ws.cell(row, col).value
            if value is not None:
                cells.append(f"{ws.cell(row, col).coordinate}={value!r}")
        if cells:
            print(f"row {row}: " + " | ".join(cells))
    wb.close()

wb = openpyxl.load_workbook(book_path, data_only=True, read_only=True)
ws = wb.worksheets[1]
print(f"dataset sheet={ws.title!r} shape={ws.max_row}x{ws.max_column}")
for row in range(1, min(ws.max_row, 24) + 1):
    cells = []
    for col in range(1, min(ws.max_column, 40) + 1):
        value = ws.cell(row, col).value
        if value is not None:
            cells.append(f"{ws.cell(row, col).coordinate}={value!r}")
    if cells:
        print(f"row {row}: " + " | ".join(cells))
wb.close()
